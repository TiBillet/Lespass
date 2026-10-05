"""
La caisse écrit UNE ligne par article, quelle que soit la façon de payer (chantier 05,
session H-1b-2).
/ The register writes ONE line per item, whatever the way of paying (chantier 05,
session H-1b-2).

LOCALISATION : tests/pytest/test_caisse_une_ligne_par_article.py

RÈGLE MÉTIER TESTÉE
Un article du panier devient UNE ligne de vente (`LigneArticle`), même quand il est payé
avec plusieurs moyens (carte NFC puis CB, jetons cadeau puis monnaie locale, deux
cartes, commande de table). Les moyens de paiement vivent seulement dans les
règlements de la vente.
- la ligne porte la vraie quantité (`qty`), le prix unitaire (`amount`) et le total
  catalogue de la formule (prix × quantité, arrondi demi-haut) ; la part payée en
  jetons cadeau (moyen LG) s'additionne dans `part_en_jetons` ;
- les lignes écrites par la caisse laissent vides le moyen, la monnaie, la carte et le
  portefeuille (Q-H2) ; elles gardent l'identifiant du paiement et le point de vente ;
- poids et volume (D15, Q-H4) : `qty` = la quantité réelle dans l'unité du prix (kg si
  le stock est en grammes ou sans stock, L si le stock est en centilitres), `amount` =
  le prix au kg / au litre ; le total vaut ce que la caisse a fait encaisser ; le poids
  saisi reste dans `weight_quantity` (le stock le lit) ; une pesée est UN article :
  une quantité de panier autre que 1 (envoi forgé) est refusée, rien n'est écrit ;
  un poids envoyé pour un tarif à la pièce (envoi forgé) est refusé de la même façon ;
- l'écran du détail, le ticket (deux imprimantes), l'admin, le rapport et son export
  tableur affichent la quantité avec son unité (« 0,350 kg ») ; deux pesées restent
  deux articles ; un produit vendu au poids et à la pièce donne une ligne par unité au
  détail du rapport ; une pesée sans prix d'achat compte pour 1 article au coût
  inconnu ; une commande de table refuse un tarif au poids ;
- retour de consigne (D13) : quantité négative, prix positif (celui du gobelet), coût
  d'achat négatif ; le rapport compte les gobelets rendus en positif ;
- vider une carte (D12) : la vente `VIDAGE_CARTE` seulement, plus aucune ligne
  « Refund » sans vente ; l'article « Jetons cadeau repris » n'a pas de moyen ;
- la garde de la cascade refuse un débit qui ne vaut pas le total de l'article ;
- l'outil de test `reset_carte` délie une carte qui a des règlements, sans erreur.
/ One cart item = ONE sale line, whatever the payment. Real quantity, unit price,
formula total, token part summed. Empty method / currency / card / wallet columns.
Weight and volume in kg / L at the price per kg / L. Deposit return: negative
quantity, positive price. Card emptying: no "Refund" line without a sale.

D'OÙ VIENNENT LES VALEURS ATTENDUES (calculées à la main)
- 3 jus à 3,50 € = 1050 ; carte 500 + CB 550 = 1050 ;
- bière 500 à 20 %, jetons 300 : HT = 300 + arrondi(200 × 100 / 120 = 166,67) = 467,
  TVA 33 ;
- comté 350 g à 12,90 €/kg : 0,350 × 1290 = 451,5 → 452 ; coût 0,350 × 800 = 280 ;
  355 g : 0,355 × 1290 = 457,95 → 458 ; deux pesées de 350 g = deux fois 452 ;
- vin 33 cl à 12,90 €/L : 0,33 × 1290 = 425,7 → 426 ; coût 0,33 × 600 = 198 ;
- gobelet 1,00 € à 20 % : −1 × 100 = −100, HT −83, TVA −17 ; coût −1 × 30 = −30 ;
- table : 2 bières à 5,00 € + 1 jus à 3,50 € = 1350 = jetons 300 + monnaie locale 1050.
/ Hand-computed values.

COMMENT CHAQUE TEST RETROUVE SA VENTE
La base est partagée avec le serveur de dev : un test ne lit que SA vente, retrouvée
par la clé d'idempotence qu'il a envoyée, par la commande de table qu'il a créée, ou
par la carte qu'il a créée (vidage). Le rapport est lu sur une période d'une
microseconde autour de l'encaissement de la vente. Chaque test finit par
`verifier_egalites(vente)` (tests/pytest/fabriques_vente.py).
/ Each test only reads its own sale (idempotency key, table order, card); reports are
read over a one-microsecond period.

SIMULATIONS
Chaque test est marqué `django_db` : la transaction est annulée à la fin
(tests/PIEGES.md 13.1), rien ne reste en base de dev. La caisse est activée EN MÉMOIRE
par `configuration_modifiee()` (tests/PIEGES.md 13.22). Celery est simulé
(tests/PIEGES.md 13.4). La cascade des monnaies du lieu écrit dans `fedow_core`, en
base locale : aucun appel réseau. Pour le vidage, le lieu est dit « non relié » à
l'ancien Fedow et le client `FedowAPI` est simulé.
/ Rolled back per test; register module on in memory; Celery faked; no network call.

Spécification : TECH_DOC/SESSIONS/COMPTABILITE/CHANTIER-05-H-retrait.md (§2, §2.1,
§5 tests 1, 2, 3, 5) ; CHANTIER-05-montants-entiers.md (§2, D12, D13, D15, D21) ;
CHANTIER-05-SUIVI.md (§4 relecture Opus de la spec H-1 ; §5 Q-H1, Q-H2, Q-H4) ;
briefs CHANTIER-05-briefs/05-H-1b-2.md et 05-H-1b-2-bis.md (tests 13 à 23),
05-H-1-ter.md (test 1).

Lancer / Run : make test ARGS="tests/pytest/test_caisse_une_ligne_par_article.py"
"""

import html
import io
import re
import uuid as uuid_module
from datetime import timedelta
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from django.http import QueryDict
from django.test import override_settings
from django.utils import timezone, translation
from django_tenants.utils import tenant_context
from openpyxl import load_workbook

from Administration.admin.site import staff_admin_site
from Administration.admin_tenant import ArticlesDeLaVenteInline, LigneArticleInline
from BaseBillet.models import (
    LigneArticle,
    Membership,
    PaymentMethod,
    Price,
    Product,
    SaleOrigin,
)
from BaseBillet.models_vente import Vente
from BaseBillet.services_refund import get_or_create_product_remboursement
from BaseBillet.services_vente import EgaliteDeVenteRompue, ouvrir_vente
from comptabilite.excel_export import generer_excel_cloture
from comptabilite.models import ClotureCaisse as ClotureCaisseUnique
from comptabilite.presentation import sections_pour_affichage
from comptabilite.rapport import RapportDesVentes
from laboutik.printing.sunmi_inner import ticket_data_to_json_commands
from laboutik.utils.test_helpers import reset_carte
from laboutik.views import _creer_lignes_articles_cascade, _extraire_articles_du_panier
from fabriques_panier import (
    configuration_modifiee,
    taches_celery_enregistrees,
)
from fabriques_vente import verifier_egalites
from inventaire.models import Stock, UniteStock
from laboutik.models import ArticleCommandeSauvegarde, CommandeSauvegarde
from laboutik.printing.escpos_builder import build_escpos_from_ticket_data
from laboutik.printing.formatters import formatter_ticket_vente
from test_caisse_ecrit_la_vente import (
    ajouter_un_solde_sur_la_carte,
    cle_de_panier,
    creer_un_article_de_caisse,
    creer_un_gobelet_et_son_retour,
    creer_une_table_occupee_avec_sa_commande,
    monnaie_cadeau_du_lieu,
    monnaie_locale_du_lieu,
    nouvelle_cle_d_idempotence,
    payer_la_commande_de_table,
    payer_le_reste_avec_une_deuxieme_carte,
    payer_le_reste_en_especes_ou_en_cb,
    payer_par_la_carte_du_client,
    payer_un_panier_compose_a_la_caisse,
    poser_les_prix_d_achat,
    reglements_de_la_vente,
    retrouver_la_vente_de_la_cle,
    seul_article_de_la_vente,
    solde_de_la_carte,
)
from test_caisse_vider_carte_deux_fedow import (
    creer_une_carte_client_sans_solde,
    monnaie_cadeau_propre_au_lieu,
    monnaie_locale_propre_au_lieu,
    poser_un_solde_sur_la_carte,
    preparer_la_caisse,
    retrouver_la_vente_de_vidage,
    vider_la_carte_a_la_caisse,
)
from test_caracterisation_caisse import (
    creer_un_point_de_vente,
    creer_une_carte_nfc_chargee,
    payer_a_la_caisse,
)
from test_caracterisation_en_ligne import creer_un_administrateur_du_lieu

pytestmark = pytest.mark.django_db

# Adresse de l'écran du détail d'une vente (laboutik/views.py, `detail_vente`).
# / Address of the sale detail screen.
DEBUT_DE_L_URL_DU_DETAIL_D_UNE_VENTE = "/laboutik/caisse/detail-vente/"


@pytest.fixture(autouse=True)
def _langue_par_defaut_apres_chaque_test():
    """Les signaux activent la langue du lieu sans la remettre (tests/PIEGES.md, 10.5).
    / Signals activate the venue language without resetting it."""
    yield
    translation.deactivate()


@pytest.fixture
def lieu(tenant):
    """
    Le lieu `lespass`, caisse activée EN MÉMOIRE, avec Celery simulé pendant tout le
    test. `configuration_modifiee()` ne sauve jamais la `Configuration`, dont le cache
    est partagé avec le serveur live (tests/PIEGES.md 13.22).
    / The `lespass` venue, register on in memory, Celery faked for the whole test.
    """
    with tenant_context(tenant):
        with configuration_modifiee(module_caisse=True, module_monnaie_locale=True):
            with taches_celery_enregistrees() as taches_demandees:
                yield SimpleNamespace(
                    tenant=tenant,
                    taches_demandees=taches_demandees,
                )


# --------------------------------------------------------------------------
# Les parcours de la caisse : chacun écrit UNE vente et la rend
# / The register paths: each one writes ONE sale and returns it
# --------------------------------------------------------------------------


def vendre_trois_jus_en_especes(lieu):
    """
    Parcours « espèces » : 3 jus à 3,50 € (TVA 20 %) payés en espèces, par la vraie
    route de paiement de la caisse (chemin à un moyen, `_creer_lignes_articles`).
    / Cash path: 3 juices paid in cash, through the real route.

    :return: objet avec `vente`, `point_de_vente`, `cle_d_idempotence`
    """
    jus = creer_un_article_de_caisse("jus", prix_en_euros="3.50", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([jus.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()

    reponse = payer_a_la_caisse(
        client_du_caissier,
        point_de_vente,
        jus,
        quantite=3,
        moyen_de_paiement="espece",
        autres_champs={"cle_idempotence_paiement": cle_d_idempotence},
    )

    assert reponse.status_code == 200
    return SimpleNamespace(
        vente=retrouver_la_vente_de_la_cle(cle_d_idempotence),
        point_de_vente=point_de_vente,
        cle_d_idempotence=cle_d_idempotence,
    )


def vendre_trois_jus_carte_500_puis_cb(lieu):
    """
    Parcours 1 : 3 jus à 3,50 € (10,50 €). La carte du client porte 5,00 € de monnaie
    locale ; le reste (5,50 €) est réglé en CB, par la route du paiement
    complémentaire (cascade, `_executer_paiement_complementaire`).
    / Path 1: 3 juices, 5.00 € on the card, the rest (5.50 €) by bank card.

    :return: objet avec `vente`, `point_de_vente`, `cle_d_idempotence`
    """
    jus = creer_un_article_de_caisse("jus", prix_en_euros="3.50", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([jus.produit])
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=500)
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()

    reponse = payer_le_reste_en_especes_ou_en_cb(
        client_du_caissier,
        point_de_vente,
        {f"repid-{cle_de_panier(jus)}": "3"},
        carte_du_client,
        moyen_du_reste="carte_bancaire",
        cle_d_idempotence=cle_d_idempotence,
    )

    assert reponse.status_code == 200
    return SimpleNamespace(
        vente=retrouver_la_vente_de_la_cle(cle_d_idempotence),
        point_de_vente=point_de_vente,
        cle_d_idempotence=cle_d_idempotence,
    )


def vendre_une_biere_300_jetons_200_monnaie_locale(lieu):
    """
    Parcours 2 : une bière à 5,00 € (TVA 20 %) payée par la carte du client seule
    (`_payer_par_nfc`) : 3,00 € de jetons cadeau, puis 2,00 € de monnaie locale.
    / Path 2: a 5.00 € beer paid by card only: 3.00 € of gift tokens, 2.00 € local.

    :return: objet avec `vente`, `point_de_vente`, `cle_d_idempotence`
    """
    biere = creer_un_article_de_caisse("biere", prix_en_euros="5.00", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([biere.produit])
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=200)
    ajouter_un_solde_sur_la_carte(carte_du_client, monnaie_cadeau_du_lieu(lieu), 300)
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()

    reponse = payer_par_la_carte_du_client(
        client_du_caissier,
        point_de_vente,
        {f"repid-{cle_de_panier(biere)}": "1"},
        carte_du_client,
        cle_d_idempotence,
    )

    assert reponse.status_code == 200
    return SimpleNamespace(
        vente=retrouver_la_vente_de_la_cle(cle_d_idempotence),
        point_de_vente=point_de_vente,
        cle_d_idempotence=cle_d_idempotence,
    )


def vendre_trois_jus_avec_deux_cartes(lieu):
    """
    Parcours 3 : 3 jus à 3,50 € (10,50 €). La carte 1 porte 5,00 € de monnaie locale,
    la carte 2 en porte 6,00 € : la carte 1 paie 5,00 €, la carte 2 paie le reste
    (5,50 €), par la route du paiement complémentaire (moyen « nfc »).
    / Path 3: card 1 pays 5.00 €, card 2 pays the rest (5.50 €).

    :return: objet avec `vente`, `point_de_vente`, `cle_d_idempotence`
    """
    jus = creer_un_article_de_caisse("jus", prix_en_euros="3.50", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([jus.produit])
    carte_1 = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=500)
    carte_2 = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=600)
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()

    reponse = payer_le_reste_avec_une_deuxieme_carte(
        client_du_caissier,
        point_de_vente,
        {f"repid-{cle_de_panier(jus)}": "3"},
        carte_1,
        carte_2,
        cle_d_idempotence,
    )

    assert reponse.status_code == 200
    return SimpleNamespace(
        vente=retrouver_la_vente_de_la_cle(cle_d_idempotence),
        point_de_vente=point_de_vente,
        cle_d_idempotence=cle_d_idempotence,
    )


def payer_une_table_en_nfc_avec_deux_monnaies(lieu):
    """
    Parcours 4 : une table occupée porte une commande de 2 bières à 5,00 € et d'un jus
    à 3,50 € (13,50 €). Le client paie avec sa carte NFC, qui porte 3,00 € de jetons
    cadeau et 20,00 € de monnaie locale (`payer_commande` → `_payer_par_nfc`). La
    cascade débite d'abord les jetons (3,00 €), puis la monnaie locale (10,50 €).
    / Path 4: a table order (2 beers + 1 juice) paid by NFC card: 3.00 € of gift tokens
    then 10.50 € of local currency.

    :return: objet avec `vente`, `point_de_vente`, `biere`, `jus`, `commande`
    """
    biere = creer_un_article_de_caisse("biere", prix_en_euros="5.00", taux_tva="20.00")
    jus = creer_un_article_de_caisse("jus", prix_en_euros="3.50", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([biere.produit, jus.produit])
    table_et_commande = creer_une_table_occupee_avec_sa_commande(biere, quantite=2)
    ArticleCommandeSauvegarde.objects.create(
        commande=table_et_commande.commande,
        product=jus.produit,
        price=jus.tarif,
        qty=1,
        statut=ArticleCommandeSauvegarde.EN_ATTENTE,
    )
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=2000)
    ajouter_un_solde_sur_la_carte(carte_du_client, monnaie_cadeau_du_lieu(lieu), 300)
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)

    reponse = payer_la_commande_de_table(
        client_du_caissier,
        point_de_vente,
        table_et_commande.commande,
        moyen_de_paiement="nfc",
        autres_champs={"tag_id": carte_du_client.tag_id},
    )

    assert reponse.status_code == 200
    commande = table_et_commande.commande
    commande.refresh_from_db()
    assert commande.statut == CommandeSauvegarde.PAID
    assert commande.vente_id is not None, "La commande payée n'a pas de vente."
    return SimpleNamespace(
        vente=Vente.objects.get(pk=commande.vente_id),
        point_de_vente=point_de_vente,
        biere=biere,
        jus=jus,
        commande=commande,
    )


def recharger_une_carte_en_especes(lieu):
    """
    Parcours « recharge » : le client recharge sa carte de 10,00 € de monnaie locale,
    payés en espèces (`_executer_recharges` → `_creer_lignes_articles`).
    / Top-up path: 10.00 € of local currency on the card, paid in cash.

    :return: objet avec `vente`, `point_de_vente`, `cle_d_idempotence`
    """
    recharge = creer_un_article_de_caisse(
        "recharge euros",
        prix_en_euros="10.00",
        taux_tva="20.00",
        methode_caisse=Product.RECHARGE_EUROS,
        asset=monnaie_locale_du_lieu(lieu),
    )
    point_de_vente = creer_un_point_de_vente([recharge.produit])
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=0)
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()

    reponse = payer_a_la_caisse(
        client_du_caissier,
        point_de_vente,
        recharge,
        quantite=1,
        moyen_de_paiement="espece",
        autres_champs={
            "tag_id": carte_du_client.tag_id,
            "cle_idempotence_paiement": cle_d_idempotence,
        },
    )

    assert reponse.status_code == 200
    return SimpleNamespace(
        vente=retrouver_la_vente_de_la_cle(cle_d_idempotence),
        point_de_vente=point_de_vente,
        cle_d_idempotence=cle_d_idempotence,
    )


# Les parcours, par leur nom : les tests paramétrés (5 et 12) les rejouent.
# / The paths by name: the parametrized tests (5 and 12) replay them.
PARCOURS_DE_LA_CAISSE = {
    "especes": vendre_trois_jus_en_especes,
    "recharge_en_especes": recharger_une_carte_en_especes,
    "carte_puis_cb": vendre_trois_jus_carte_500_puis_cb,
    "jetons_puis_monnaie_locale": vendre_une_biere_300_jetons_200_monnaie_locale,
    "deux_cartes": vendre_trois_jus_avec_deux_cartes,
    "table_en_nfc_deux_monnaies": payer_une_table_en_nfc_avec_deux_monnaies,
}


# --------------------------------------------------------------------------
# Assistants de lecture
# / Reading helpers
# --------------------------------------------------------------------------


def article_du_tarif(vente, tarif):
    """
    L'UNIQUE ligne de la vente pour ce tarif, relue en base. Deux lignes pour un même
    tarif veulent dire que l'article a été coupé en parts.
    / The sale's ONLY line for this price. Two lines mean the item was split in parts.
    """
    lignes_du_tarif = list(vente.articles.filter(pricesold__price=tarif))
    assert len(lignes_du_tarif) == 1, (
        f"Attendu : une seule ligne pour le tarif {tarif.name}, "
        f"trouvé : {len(lignes_du_tarif)}."
    )
    return lignes_du_tarif[0]


def champ_vide(valeur):
    """
    Vrai si un champ de la ligne est vide : None (clé étrangère, uuid, texte nul) ou
    texte vide.
    / True when a line field is empty: None or an empty text.
    """
    return valeur is None or valeur == ""


def texte_normalise(texte):
    """
    Le texte sans balises HTML, entités décodées, espaces insécables remplacées par
    des espaces ordinaires, espaces multiples réduites à une seule.
    / Text without HTML tags, entities decoded, non-breaking spaces made plain.
    """
    texte_sans_balises = re.sub(r"<[^>]+>", " ", texte)
    texte_decode = html.unescape(texte_sans_balises)
    texte_sans_insecables = texte_decode.replace("\xa0", " ").replace(" ", " ")
    return re.sub(r"\s+", " ", texte_sans_insecables).strip()


def textes_des_cellules(contenu_html, testid):
    """
    Les textes (normalisés) de toutes les cellules qui portent ce `data-testid`.
    / The normalised texts of every cell carrying this data-testid.
    """
    motif_de_la_cellule = re.compile(
        r'<td[^>]*data-testid="' + re.escape(testid) + r'"[^>]*>(.*?)</td>',
        re.DOTALL,
    )
    textes = []
    for contenu_de_la_cellule in motif_de_la_cellule.findall(contenu_html):
        textes.append(texte_normalise(contenu_de_la_cellule))
    return textes


def rapport_de_l_instant_de_la_vente(vente):
    """
    Le rapport des ventes sur la seule microseconde de l'encaissement de la vente :
    il ne lit que cette vente.
    / The sales report over the sale's settlement microsecond only.
    """
    vente_relue = Vente.objects.get(pk=vente.pk)
    instant_de_l_encaissement = vente_relue.datetime_encaissement
    return RapportDesVentes(
        debut=instant_de_l_encaissement,
        fin=instant_de_l_encaissement + timedelta(microseconds=1),
    )


def quantite_et_unite(texte_de_la_quantite):
    """
    Coupe « 0,350 kg » en (Decimal("0.350"), "kg"). Rend (None, texte) si le texte
    n'a pas la forme « nombre unité ».
    / Splits "0,350 kg" into (Decimal("0.350"), "kg").
    """
    correspondance = re.fullmatch(
        r"(-?[0-9]+(?:,[0-9]+)?) ?([A-Za-z]+)", texte_de_la_quantite
    )
    if correspondance is None:
        return (None, texte_de_la_quantite)
    nombre = Decimal(correspondance.group(1).replace(",", "."))
    return (nombre, correspondance.group(2))


# --------------------------------------------------------------------------
# Ce qu'il faut pour vendre au poids ou au volume
# / What weight and volume sales need
# --------------------------------------------------------------------------


def creer_un_comte_au_kilo(prix_achat_en_centimes=800, avec_stock_en_grammes=False):
    """
    Le comté vendu au poids : 12,90 €/kg (TVA 5,5 %), prix d'achat au kilo.
    Sans stock, la caisse compte en grammes (`_diviseur_de_la_quantite_saisie`).
    / Cheese sold by weight: 12.90 €/kg; without stock, the register counts grams.

    :param prix_achat_en_centimes: prix d'achat au kilo ; 0 = inconnu (D21)
    :param avec_stock_en_grammes: True pour lier un stock en grammes au produit
    """
    comte = creer_un_article_de_caisse(
        "comte au poids",
        prix_en_euros="12.90",
        taux_tva="5.50",
        prix_achat_en_centimes=prix_achat_en_centimes,
        vendu_au_poids=True,
    )
    if avec_stock_en_grammes:
        Stock.objects.create(
            product=comte.produit,
            quantite=10000,
            unite=UniteStock.GR,
        )
    return comte


def champs_d_une_pesee(article, quantite_saisie, montant_en_centimes):
    """
    Les champs du panier pour UNE pesée, comme addition.js les envoie : la clé porte
    le suffixe `--1` (une pesée = une ligne de panier, tests/PIEGES.md 66), la
    quantité de panier vaut 1, `weight-` porte la quantité saisie (g ou cl),
    `custom-` le montant calculé par l'écran.
    / The cart fields of ONE weighing, like addition.js sends them.
    """
    cle_de_la_ligne = cle_de_panier(article, numero_de_ligne=1)
    return {
        f"repid-{cle_de_la_ligne}": "1",
        f"weight-{cle_de_la_ligne}": str(quantite_saisie),
        f"custom-{cle_de_la_ligne}": str(montant_en_centimes),
    }


def vendre_une_pesee(
    lieu, article, quantite_saisie, montant_en_centimes, facon_de_payer
):
    """
    Vend UNE pesée par la vraie route de la caisse.
    / Sells ONE weighing through the real register route.

    `facon_de_payer` :
    - « espece » : chemin à un moyen (`_creer_lignes_articles`) ;
    - « nfc » : la carte du client, monnaie locale seule (cascade) ;
    - « nfc_deux_monnaies » : la carte du client, 2,00 € de jetons cadeau puis la
      monnaie locale (cascade, deux débits sur le même article).

    :return: objet avec `vente`, `cle_d_idempotence`
    """
    point_de_vente = creer_un_point_de_vente([article.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()
    champs_du_panier = champs_d_une_pesee(article, quantite_saisie, montant_en_centimes)

    if facon_de_payer == "espece":
        reponse = payer_un_panier_compose_a_la_caisse(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            moyen_de_paiement="espece",
            cle_d_idempotence=cle_d_idempotence,
        )
    elif facon_de_payer == "nfc":
        carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=2000)
        reponse = payer_par_la_carte_du_client(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            carte_du_client,
            cle_d_idempotence,
        )
    elif facon_de_payer == "nfc_deux_monnaies":
        carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=2000)
        ajouter_un_solde_sur_la_carte(
            carte_du_client, monnaie_cadeau_du_lieu(lieu), 200
        )
        reponse = payer_par_la_carte_du_client(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            carte_du_client,
            cle_d_idempotence,
        )
    else:
        raise ValueError(f"Façon de payer inconnue : {facon_de_payer}")

    assert reponse.status_code == 200, reponse.content.decode()[:400]
    return SimpleNamespace(
        vente=retrouver_la_vente_de_la_cle(cle_d_idempotence),
        cle_d_idempotence=cle_d_idempotence,
    )


# --------------------------------------------------------------------------
# 1 — Carte puis CB : trois jus, UNE ligne de quantité 3
# / 1 — Card then bank card: three juices, ONE line of quantity 3
# --------------------------------------------------------------------------


def test_nfc_trois_jus_une_seule_ligne_qty_3(lieu):
    """
    3 jus à 3,50 € payés 5,00 € par la carte (monnaie locale) + 5,50 € en CB
    (paiement complémentaire). UNE ligne : quantité 3, prix unitaire 350, total 1050.
    Deux règlements : monnaie locale 500, CB 550.
    / 3 juices paid 5.00 € by card + 5.50 € by bank card: ONE line (3, 350, 1050);
    payments LE 500 + CB 550.
    """
    parcours = vendre_trois_jus_carte_500_puis_cb(lieu)
    vente = parcours.vente

    article = seul_article_de_la_vente(vente)
    assert article.qty == Decimal("3")
    assert article.amount == 350
    assert article.total_catalogue == 1050
    assert article.part_offerte == 0
    assert article.part_en_jetons == 0
    assert article.total_ttc == 1050
    assert article.total_ht == 875
    assert article.total_tva == 175
    assert reglements_de_la_vente(vente) == [
        (PaymentMethod.CC, 550),
        (PaymentMethod.LOCAL_EURO, 500),
    ]
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 2 — Jetons puis monnaie locale : UNE ligne, part en jetons 300
# / 2 — Tokens then local currency: ONE line, token part 300
# --------------------------------------------------------------------------


def test_jetons_une_ligne_part_en_jetons_300(lieu):
    """
    Une bière à 5,00 € (TVA 20 %) payée par la carte seule : 3,00 € de jetons cadeau
    puis 2,00 € de monnaie locale. UNE ligne : quantité 1, prix 500, total 500, rien
    d'offert, part payée en jetons 300 ; la TVA porte sur le reste (200) : HT 467,
    TVA 33. Règlements : jetons 300, monnaie locale 200.
    / A 5.00 € beer paid 3.00 € in tokens + 2.00 € local: ONE line, token part 300,
    HT 467, VAT 33; payments LG 300 + LE 200.
    """
    parcours = vendre_une_biere_300_jetons_200_monnaie_locale(lieu)
    vente = parcours.vente

    article = seul_article_de_la_vente(vente)
    assert article.qty == Decimal("1")
    assert article.amount == 500
    assert article.total_catalogue == 500
    assert article.part_offerte == 0
    assert article.source_offert == ""
    assert article.part_en_jetons == 300
    assert article.total_ttc == 500
    assert article.total_ht == 467
    assert article.total_tva == 33
    assert reglements_de_la_vente(vente) == [
        (PaymentMethod.LOCAL_EURO, 200),
        (PaymentMethod.LOCAL_GIFT, 300),
    ]
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 3 — Deux cartes : UNE ligne, deux règlements
# / 3 — Two cards: ONE line, two payments
# --------------------------------------------------------------------------


def test_deuxieme_carte_une_ligne(lieu):
    """
    3 jus à 3,50 € payés par deux cartes : 5,00 € sur la carte 1, 5,50 € sur la
    carte 2. UNE ligne (quantité 3, prix 350, total 1050) ; deux règlements en
    monnaie locale, 500 et 550.
    / 3 juices paid by two cards: ONE line (3, 350, 1050); two LE payments 500 + 550.
    """
    parcours = vendre_trois_jus_avec_deux_cartes(lieu)
    vente = parcours.vente

    article = seul_article_de_la_vente(vente)
    assert article.qty == Decimal("3")
    assert article.amount == 350
    assert article.total_catalogue == 1050
    assert article.total_ttc == 1050
    assert reglements_de_la_vente(vente) == [
        (PaymentMethod.LOCAL_EURO, 500),
        (PaymentMethod.LOCAL_EURO, 550),
    ]
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 4 — Commande de table payée en NFC avec deux monnaies : une ligne par article
# / 4 — Table order paid by NFC with two currencies: one line per item
# --------------------------------------------------------------------------


def test_table_payee_en_nfc_une_ligne_par_article(lieu):
    """
    La commande d'une table (2 bières à 5,00 € + 1 jus à 3,50 €) est payée par la
    carte du client : 3,00 € de jetons cadeau puis 10,50 € de monnaie locale.
    Deux lignes, une par article : bières (quantité 2, prix 500, total 1000) et jus
    (quantité 1, prix 350, total 350). Les jetons sont dans la part en jetons des
    lignes (300 en tout). Règlements : jetons 300, monnaie locale 1050.
    / A table order paid by card (tokens 300 + local 1050): two lines, one per item,
    with the real quantities; token parts add up to 300.
    """
    parcours = payer_une_table_en_nfc_avec_deux_monnaies(lieu)
    vente = parcours.vente

    assert vente.articles.count() == 2
    ligne_des_bieres = article_du_tarif(vente, parcours.biere.tarif)
    assert ligne_des_bieres.qty == Decimal("2")
    assert ligne_des_bieres.amount == 500
    assert ligne_des_bieres.total_catalogue == 1000

    ligne_du_jus = article_du_tarif(vente, parcours.jus.tarif)
    assert ligne_du_jus.qty == Decimal("1")
    assert ligne_du_jus.amount == 350
    assert ligne_du_jus.total_catalogue == 350

    somme_des_parts_en_jetons = 0
    for ligne in vente.articles.all():
        somme_des_parts_en_jetons += ligne.part_en_jetons
    assert somme_des_parts_en_jetons == 300

    assert reglements_de_la_vente(vente) == [
        (PaymentMethod.LOCAL_EURO, 1050),
        (PaymentMethod.LOCAL_GIFT, 300),
    ]
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 5 — Les lignes de la caisse n'ont ni moyen, ni monnaie, ni carte, ni portefeuille
# / 5 — Register lines have no method, currency, card or wallet
# --------------------------------------------------------------------------


@pytest.mark.parametrize("nom_du_parcours", list(PARCOURS_DE_LA_CAISSE.keys()))
def test_lignes_de_caisse_sans_moyen_ni_monnaie_ni_carte(lieu, nom_du_parcours):
    """
    Pour chaque parcours de la caisse (espèces, recharge en espèces, carte puis CB,
    jetons puis monnaie locale, deux cartes, table en NFC) : chaque ligne écrite
    laisse vides le moyen
    (`payment_method`), la monnaie (`asset`), la carte et le portefeuille (Q-H2). Les
    règlements portent ces informations. La ligne garde l'identifiant du paiement
    (`uuid_transaction`, idempotence de la caisse) et le point de vente.
    / For each register path: every line leaves method, currency, card and wallet
    empty; it keeps the payment id and the point of sale.
    """
    parcours = PARCOURS_DE_LA_CAISSE[nom_du_parcours](lieu)
    vente = parcours.vente

    lignes_de_la_vente = list(vente.articles.all())
    assert len(lignes_de_la_vente) >= 1
    for ligne in lignes_de_la_vente:
        assert champ_vide(ligne.payment_method), (
            f"Moyen de la ligne : {ligne.payment_method!r} (attendu : vide)."
        )
        assert champ_vide(ligne.asset), f"Monnaie de la ligne : {ligne.asset!r}."
        assert ligne.carte_id is None, f"Carte de la ligne : {ligne.carte_id!r}."
        assert ligne.wallet_id is None, (
            f"Portefeuille de la ligne : {ligne.wallet_id!r}."
        )
        assert ligne.uuid_transaction is not None
        assert ligne.point_de_vente_id == parcours.point_de_vente.pk
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 6 — Poids : UNE ligne de 0,350 kg au prix du kilo
# / 6 — Weight: ONE line of 0.350 kg at the price per kilo
# --------------------------------------------------------------------------


@pytest.mark.parametrize("facon_de_payer", ["espece", "nfc", "nfc_deux_monnaies"])
def test_fromage_une_ligne_0_350_kg_prix_au_kilo(lieu, facon_de_payer):
    """
    350 g de comté à 12,90 €/kg (prix d'achat 8,00 €/kg), sans stock (donc en
    grammes) : UNE ligne, quantité 0,350 (kg), prix unitaire 1290 (le prix du kilo),
    total 452 (ce que la caisse fait encaisser : 0,350 × 1290 = 451,5 → 452), coût
    d'achat sur 0,350 kg (280). Le poids saisi reste dans `weight_quantity` (350).
    Même chose en espèces, par la carte (monnaie locale seule) et par la carte avec
    2,00 € de jetons cadeau (part en jetons 200).
    / 350 g of cheese at 12.90 €/kg: ONE line, qty 0.350, price 1290, total 452, cost
    280; same in cash, by card, and by card with 2.00 € of tokens.
    """
    comte = creer_un_comte_au_kilo(prix_achat_en_centimes=800)

    parcours = vendre_une_pesee(
        lieu,
        comte,
        quantite_saisie=350,
        montant_en_centimes=452,
        facon_de_payer=facon_de_payer,
    )
    vente = parcours.vente

    article = seul_article_de_la_vente(vente)
    assert article.qty == Decimal("0.350")
    assert article.amount == 1290
    assert article.total_catalogue == 452
    assert article.cout_achat == 280
    assert article.weight_quantity == 350

    somme_des_reglements = 0
    for reglement in vente.reglements.all():
        somme_des_reglements += reglement.montant
    assert somme_des_reglements == 452

    if facon_de_payer == "nfc_deux_monnaies":
        assert article.part_en_jetons == 200
        assert reglements_de_la_vente(vente) == [
            (PaymentMethod.LOCAL_EURO, 252),
            (PaymentMethod.LOCAL_GIFT, 200),
        ]
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 6 bis — Une pesée avec une quantité de panier 2 : refusée, rien n'est écrit
# / 6 bis — A weighing with a cart quantity of 2: refused, nothing written
# --------------------------------------------------------------------------


@pytest.mark.parametrize("facon_de_payer", ["espece", "nfc", "carte_puis_cb"])
def test_pesee_en_quantite_2_refusee_rien_n_est_ecrit(lieu, facon_de_payer):
    """
    Un envoi forgé porte une pesée (350 g de comté) avec une quantité de panier 2.
    L'écran ne l'envoie jamais : chaque pesée a sa propre ligne de panier, en
    quantité 1 (tests/PIEGES.md 66). La caisse refuse (400), avant toute écriture :
    ni vente, ni ligne, ni débit de la carte. Même refus en espèces, par la carte du
    client, et par le paiement complémentaire (carte puis CB).
    / A forged post with a weighing in quantity 2 is refused (400): no sale, no line,
    no card debit; in cash, by card and by the complement payment.
    """
    comte = creer_un_comte_au_kilo()
    point_de_vente = creer_un_point_de_vente([comte.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=2000)
    cle_d_idempotence = nouvelle_cle_d_idempotence()
    champs_du_panier = champs_d_une_pesee(
        comte, quantite_saisie=350, montant_en_centimes=452
    )
    cle_de_la_ligne = cle_de_panier(comte, numero_de_ligne=1)
    champs_du_panier[f"repid-{cle_de_la_ligne}"] = "2"

    if facon_de_payer == "espece":
        reponse = payer_un_panier_compose_a_la_caisse(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            moyen_de_paiement="espece",
            cle_d_idempotence=cle_d_idempotence,
        )
    elif facon_de_payer == "nfc":
        reponse = payer_par_la_carte_du_client(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            carte_du_client,
            cle_d_idempotence,
        )
    else:
        reponse = payer_le_reste_en_especes_ou_en_cb(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            carte_du_client,
            moyen_du_reste="carte_bancaire",
            cle_d_idempotence=cle_d_idempotence,
        )

    assert reponse.status_code == 400, reponse.content.decode()[:400]
    assert not Vente.objects.filter(idempotency_key=cle_d_idempotence).exists()
    assert not LigneArticle.objects.filter(pricesold__price=comte.tarif).exists()
    assert solde_de_la_carte(carte_du_client, monnaie_locale_du_lieu(lieu)) == 2000


# --------------------------------------------------------------------------
# 6 ter — Un poids forgé sur un tarif à la pièce : refusé, rien n'est écrit
# / 6 ter — A forged weight on a per-item price: refused, nothing written
# --------------------------------------------------------------------------


@pytest.mark.parametrize("facon_de_payer", ["espece", "nfc"])
def test_poids_forge_sur_un_tarif_a_la_piece_refuse(lieu, facon_de_payer):
    """
    Un envoi forgé porte 3 jus (tarif à la pièce, pas au poids) avec un poids de
    350 (`weight-<jus>`). L'écran n'envoie jamais de poids pour un tarif à la pièce.
    La caisse refuse (400), comme une pesée en quantité 2, avant toute écriture : ni
    vente, ni ligne, ni débit de la carte. Sans ce refus, la ligne garderait le poids
    et passerait pour une pesée (« kg » au ticket et à l'admin, avoir partiel refusé).
    / A forged post sends a weight for a per-item price: refused (400) before writing
    anything, like a weighing in quantity 2; in cash and by card.
    """
    jus = creer_un_article_de_caisse("jus", prix_en_euros="3.50", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([jus.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=2000)
    cle_d_idempotence = nouvelle_cle_d_idempotence()
    champs_du_panier = {
        f"repid-{cle_de_panier(jus)}": "3",
        f"weight-{cle_de_panier(jus)}": "350",
    }

    if facon_de_payer == "espece":
        reponse = payer_un_panier_compose_a_la_caisse(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            moyen_de_paiement="espece",
            cle_d_idempotence=cle_d_idempotence,
        )
    else:
        reponse = payer_par_la_carte_du_client(
            client_du_caissier,
            point_de_vente,
            champs_du_panier,
            carte_du_client,
            cle_d_idempotence,
        )

    assert reponse.status_code == 400, reponse.content.decode()[:400]
    assert not Vente.objects.filter(idempotency_key=cle_d_idempotence).exists()
    assert not LigneArticle.objects.filter(pricesold__price=jus.tarif).exists()
    assert solde_de_la_carte(carte_du_client, monnaie_locale_du_lieu(lieu)) == 2000


# --------------------------------------------------------------------------
# 7 — Volume : stock en centilitres, quantité en litres
# / 7 — Volume: stock in centilitres, quantity in litres
# --------------------------------------------------------------------------


@pytest.mark.parametrize("facon_de_payer", ["espece", "nfc"])
def test_vin_au_volume_quantite_en_litres(lieu, facon_de_payer):
    """
    33 cl de vin à 12,90 €/L (prix d'achat 6,00 €/L), stock en centilitres : UNE
    ligne, quantité 0,33 (L), prix unitaire 1290 (le prix du litre), total 426
    (0,33 × 1290 = 425,7 → 426), coût d'achat sur 0,33 L (198). Le volume saisi reste
    dans `weight_quantity` (33).
    / 33 cl of wine at 12.90 €/L, stock in cl: ONE line, qty 0.33 L, price 1290,
    total 426, cost 198.
    """
    vin = creer_un_article_de_caisse(
        "vin au volume",
        prix_en_euros="12.90",
        taux_tva="20.00",
        prix_achat_en_centimes=600,
        vendu_au_poids=True,
    )
    Stock.objects.create(product=vin.produit, quantite=10000, unite=UniteStock.CL)

    parcours = vendre_une_pesee(
        lieu,
        vin,
        quantite_saisie=33,
        montant_en_centimes=426,
        facon_de_payer=facon_de_payer,
    )
    vente = parcours.vente

    article = seul_article_de_la_vente(vente)
    assert article.qty == Decimal("0.33")
    assert article.amount == 1290
    assert article.total_catalogue == 426
    assert article.cout_achat == 198
    assert article.weight_quantity == 33
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 8 — Écran du détail et ticket : « 0,350 kg » au prix du kilo
# / 8 — Detail screen and receipt: "0.350 kg" at the price per kilo
# --------------------------------------------------------------------------


@pytest.mark.parametrize("avec_stock_en_grammes", [True, False])
def test_affichage_et_ticket_du_poids(lieu, avec_stock_en_grammes):
    """
    350 g de comté à 12,90 €/kg, payés en espèces, avec un stock en grammes ou sans
    stock (la caisse compte alors en grammes : même règle que le prix).
    - l'écran du détail de la vente montre la quantité « 0,350 kg » (pas un poids
      recalculé en grammes), le prix « 12,90 €/kg » et le total « 4,52 € » ;
    - le ticket imprimé montre « 0,350 kg », le prix au kilo et le total 4,52.
    / The detail screen shows "0,350 kg", "12,90 €/kg", "4,52 €"; the receipt shows
    "0,350 kg", the price per kilo and 4.52; with a gram stock or without stock.
    """
    comte = creer_un_comte_au_kilo(avec_stock_en_grammes=avec_stock_en_grammes)
    parcours = vendre_une_pesee(
        lieu,
        comte,
        quantite_saisie=350,
        montant_en_centimes=452,
        facon_de_payer="espece",
    )
    vente = parcours.vente
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)

    # L'écran du détail de la vente.
    # / The sale detail screen.
    reponse = client_du_caissier.get(
        f"{DEBUT_DE_L_URL_DU_DETAIL_D_UNE_VENTE}{vente.uuid}/"
    )
    contenu_du_detail = reponse.content.decode()
    assert reponse.status_code == 200, contenu_du_detail[:400]
    assert textes_des_cellules(contenu_du_detail, "detail-qty") == ["0,350 kg"]
    assert textes_des_cellules(contenu_du_detail, "detail-prix-unit") == ["12,90 €/kg"]
    assert textes_des_cellules(contenu_du_detail, "detail-total-ligne") == ["4,52 €"]

    # Le ticket, tel que l'imprimante ESC/POS l'imprimerait.
    # / The receipt, as the ESC/POS printer would print it.
    donnees_du_ticket = formatter_ticket_vente(vente, None)
    texte_du_ticket = build_escpos_from_ticket_data(
        576, dict(donnees_du_ticket)
    ).decode("utf-8", errors="ignore")
    texte_du_ticket = texte_normalise(texte_du_ticket)
    assert "0,350 kg" in texte_du_ticket, texte_du_ticket
    assert "12,90" in texte_du_ticket and "/kg" in texte_du_ticket, texte_du_ticket
    assert "4.52" in texte_du_ticket or "4,52" in texte_du_ticket, texte_du_ticket
    assert "350g" not in texte_du_ticket, texte_du_ticket
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 9 — Rapport : quantité en kg, une pesée au coût inconnu compte pour 1 article
# / 9 — Report: quantity in kg, a weighing with unknown cost counts as 1 item
# --------------------------------------------------------------------------


def test_rapport_quantites_en_kg_et_articles_au_cout_inconnu(lieu):
    """
    350 g de comté à 12,90 €/kg (stock en grammes), SANS prix d'achat (0 = inconnu),
    payés en espèces. Le rapport de cette seule vente :
    - détail des ventes par produit : la quantité du comté est 0,350 avec l'unité
      « kg » (pas « 1 », le nombre de pesées) ;
    - marge brute : la pesée sans prix d'achat compte pour 1 article au coût inconnu
      (une pesée = un article, pas 0,350 article arrondi à 0).
    / Report of this single sale: comté quantity 0.350 "kg"; the weighing without
    purchase price counts as 1 item with unknown cost.
    """
    comte = creer_un_comte_au_kilo(prix_achat_en_centimes=0, avec_stock_en_grammes=True)
    parcours = vendre_une_pesee(
        lieu,
        comte,
        quantite_saisie=350,
        montant_en_centimes=452,
        facon_de_payer="espece",
    )
    vente = parcours.vente
    rapport = rapport_de_l_instant_de_la_vente(vente)

    # Le détail affiché : la ligne du comté dans « Ventes par produit ».
    # / The displayed detail: the comté row of "Sales by product".
    sections_affichees = sections_pour_affichage(
        {
            "chiffre_affaires": rapport.section_chiffre_affaires(),
            "detail": rapport.section_detail(),
        }
    )
    tableau_par_produit = None
    for section in sections_affichees:
        for tableau in section["tableaux"]:
            if tableau["testid"] == "detail-ventes-par-produit":
                tableau_par_produit = tableau
    assert tableau_par_produit is not None, "Le tableau « Ventes par produit » manque."
    lignes_du_comte = []
    for ligne_du_tableau in tableau_par_produit["lignes"]:
        if ligne_du_tableau[1]["texte"] == comte.produit.name:
            lignes_du_comte.append(ligne_du_tableau)
    assert len(lignes_du_comte) == 1
    cellule_de_la_quantite = lignes_du_comte[0][2]
    quantite_lue, unite_lue = quantite_et_unite(
        texte_normalise(cellule_de_la_quantite["texte"])
    )
    assert quantite_lue == Decimal("0.350"), cellule_de_la_quantite["texte"]
    assert unite_lue == "kg", cellule_de_la_quantite["texte"]

    # La marge brute : un article au coût inconnu.
    # / Gross margin: one item with unknown cost.
    marge_brute = rapport.section_marge_brute()
    assert marge_brute["nombre_d_articles_au_cout_inconnu"] == 1
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 10 — Retour de consigne : quantité négative, prix positif
# / 10 — Deposit return: negative quantity, positive price
# --------------------------------------------------------------------------


@pytest.mark.parametrize("moyen_de_paiement", ["espece", "nfc"])
def test_retour_consigne_quantite_negative_prix_positif(lieu, moyen_de_paiement):
    """
    Le client rapporte un gobelet consigné (1,00 €, TVA 20 %, coûté 0,30 € au lieu).
    La caisse lui rend 1,00 € (espèces, ou recrédit de sa carte). La vente est un
    AVOIR ; sa ligne a une quantité −1, le prix du gobelet en POSITIF (100), un total
    −100 (HT −83, TVA −17) et un coût d'achat négatif (−30 : le gobelet rendu retire
    son coût de la marge). Le rapport compte 1 gobelet rendu (un nombre positif).
    / A returned cup: CREDIT NOTE; line qty −1, price +100, total −100, cost −30; the
    report counts 1 returned cup.
    """
    consigne = creer_un_gobelet_et_son_retour(lieu)
    poser_les_prix_d_achat(consigne, prix_achat_du_gobelet=30, prix_achat_du_retour=99)
    point_de_vente = creer_un_point_de_vente([consigne.retour.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()
    autres_champs = {"cle_idempotence_paiement": cle_d_idempotence}
    if moyen_de_paiement == "nfc":
        carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=0)
        autres_champs["tag_id"] = carte_du_client.tag_id

    reponse = payer_a_la_caisse(
        client_du_caissier,
        point_de_vente,
        consigne.retour,
        quantite=1,
        moyen_de_paiement=moyen_de_paiement,
        autres_champs=autres_champs,
    )

    assert reponse.status_code == 200
    vente = retrouver_la_vente_de_la_cle(cle_d_idempotence)
    assert vente.nature == Vente.Nature.AVOIR
    article = seul_article_de_la_vente(vente)
    assert article.qty == Decimal("-1")
    assert article.amount == 100
    assert article.vat == Decimal("20")
    assert article.total_catalogue == -100
    assert article.total_ttc == -100
    assert article.total_ht == -83
    assert article.total_tva == -17
    assert article.cout_achat == -30

    rapport = rapport_de_l_instant_de_la_vente(vente)
    retours_de_consigne = rapport.section_annexe()["avoirs"]["retours_consigne"]
    assert retours_de_consigne["nombre"] == 1
    assert retours_de_consigne["total_en_centimes"] == -100
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 11 — Vider une carte : la vente VIDAGE_CARTE, aucune ligne « Refund »
# / 11 — Emptying a card: the VIDAGE_CARTE sale, no "Refund" line
# --------------------------------------------------------------------------


def test_vidage_sans_ligne_refund(lieu):
    """
    Le lieu n'est pas relié à l'ancien Fedow. La carte du client porte 5,00 € de
    monnaie locale du lieu ; le caissier la vide. La vente `VIDAGE_CARTE` est écrite
    (règlements monnaie locale +500, espèces −500). Aucune ligne « Refund » sans vente
    n'est écrite : ni sur la carte du client, ni pour le produit « Refund », pendant
    le test (D12 : la vente remplace ces lignes).
    / Emptying a card writes the VIDAGE_CARTE sale and no "Refund" line without a sale.
    """
    instant_du_debut_du_test = timezone.now()
    caisse = preparer_la_caisse(lieu)
    carte_du_client = creer_une_carte_client_sans_solde(lieu)
    poser_un_solde_sur_la_carte(
        carte_du_client, monnaie_locale_propre_au_lieu(lieu), 500
    )

    reponse, _classe_fedow_api = vider_la_carte_a_la_caisse(
        caisse, carte_du_client, MagicMock(), lieu_relie=False
    )

    assert reponse.status_code == 200
    vente = retrouver_la_vente_de_vidage(carte_du_client)
    assert vente.statut == Vente.Statut.REGLEE
    assert reglements_de_la_vente(vente) == [
        (PaymentMethod.CASH, -500),
        (PaymentMethod.LOCAL_EURO, 500),
    ]

    lignes_sans_vente_de_la_carte = LigneArticle.objects.filter(
        carte=carte_du_client, vente__isnull=True
    )
    assert lignes_sans_vente_de_la_carte.count() == 0

    produit_refund = get_or_create_product_remboursement()
    lignes_refund_ecrites_pendant_le_test = LigneArticle.objects.filter(
        pricesold__productsold__product=produit_refund,
        vente__isnull=True,
        datetime__gte=instant_du_debut_du_test,
    )
    assert lignes_refund_ecrites_pendant_le_test.count() == 0
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 12 — Garde : aucune quantité fractionnaire à la caisse, hors poids et volume
# / 12 — Guard: no fractional quantity at the register, except weight and volume
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "nom_du_parcours",
    [
        "carte_puis_cb",
        "jetons_puis_monnaie_locale",
        "deux_cartes",
        "table_en_nfc_deux_monnaies",
    ],
)
def test_garde_aucune_ligne_fractionnaire_a_la_caisse(lieu, nom_du_parcours):
    """
    Après chacun des parcours 1 à 4 (paiements à plusieurs moyens), aucune ligne de
    la vente n'a une quantité non entière, hors vente au poids ou au volume (la ligne
    porte alors un poids dans `weight_quantity`). Une quantité fractionnaire veut dire
    qu'un article a été coupé en parts.
    / After paths 1 to 4, no line has a non-whole quantity, except weight / volume.
    """
    parcours = PARCOURS_DE_LA_CAISSE[nom_du_parcours](lieu)
    vente = parcours.vente

    quantites_fractionnaires = []
    for ligne in vente.articles.all():
        ligne_au_poids_ou_au_volume = bool(ligne.weight_quantity)
        if ligne_au_poids_ou_au_volume:
            continue
        quantite_entiere = ligne.qty == ligne.qty.to_integral_value()
        if not quantite_entiere:
            quantites_fractionnaires.append(ligne.qty)
    assert quantites_fractionnaires == [], (
        f"Quantités fractionnaires écrites par la caisse : {quantites_fractionnaires}."
    )
    verifier_egalites(vente)


# ==========================================================================
# Session H-1b-2-bis : corrections de la relecture (tests 13 à 23)
# / Session H-1b-2-bis: review fixes (tests 13 to 23)
# ==========================================================================


def texte_imprime_par_l_imprimante_escpos(donnees_du_ticket):
    """
    Le texte que l'imprimante ESC/POS imprimerait pour ce ticket, normalisé.
    / The text the ESC/POS printer would print, normalised.
    """
    octets_imprimes = build_escpos_from_ticket_data(576, dict(donnees_du_ticket))
    return texte_normalise(octets_imprimes.decode("utf-8", errors="ignore"))


def texte_imprime_par_l_imprimante_sunmi(donnees_du_ticket):
    """
    Les textes envoyés à l'imprimante Sunmi intégrée, mis bout à bout et normalisés.
    / The texts sent to the built-in Sunmi printer, joined and normalised.
    """
    textes = []
    for commande in ticket_data_to_json_commands(dict(donnees_du_ticket)):
        textes.append(str(commande.get("value", "")))
    return texte_normalise(" ".join(textes))


# --------------------------------------------------------------------------
# 13 — Deux pesées : deux articles au ticket et au détail
# / 13 — Two weighings: two items on the receipt and the detail
# --------------------------------------------------------------------------


def test_deux_pesees_deux_articles_au_ticket(lieu):
    """
    Deux pesées de 350 g de comté à 12,90 €/kg dans le même panier, en espèces : deux
    lignes de 0,350 kg (4,52 € chacune). L'écran du détail et le ticket montrent DEUX
    articles « 0,350 kg × 12,90 €/kg 4,52 », jamais un article « 0,700 kg … 9,04 ».
    / Two 350 g weighings: the detail screen and the receipt show TWO items, never one
    "0,700 kg" item.
    """
    comte = creer_un_comte_au_kilo(avec_stock_en_grammes=True)
    point_de_vente = creer_un_point_de_vente([comte.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()
    champs_du_panier = {}
    for numero_de_la_pesee in [1, 2]:
        cle_de_la_ligne = cle_de_panier(comte, numero_de_ligne=numero_de_la_pesee)
        champs_du_panier[f"repid-{cle_de_la_ligne}"] = "1"
        champs_du_panier[f"weight-{cle_de_la_ligne}"] = "350"
        champs_du_panier[f"custom-{cle_de_la_ligne}"] = "452"

    reponse = payer_un_panier_compose_a_la_caisse(
        client_du_caissier,
        point_de_vente,
        champs_du_panier,
        moyen_de_paiement="espece",
        cle_d_idempotence=cle_d_idempotence,
    )

    assert reponse.status_code == 200, reponse.content.decode()[:400]
    vente = retrouver_la_vente_de_la_cle(cle_d_idempotence)
    assert vente.articles.count() == 2

    # L'écran du détail de la vente.
    # / The sale detail screen.
    reponse_du_detail = client_du_caissier.get(
        f"{DEBUT_DE_L_URL_DU_DETAIL_D_UNE_VENTE}{vente.uuid}/"
    )
    contenu_du_detail = reponse_du_detail.content.decode()
    assert reponse_du_detail.status_code == 200, contenu_du_detail[:400]
    assert textes_des_cellules(contenu_du_detail, "detail-qty") == [
        "0,350 kg",
        "0,350 kg",
    ]
    assert textes_des_cellules(contenu_du_detail, "detail-total-ligne") == [
        "4,52 €",
        "4,52 €",
    ]

    # Le ticket : deux articles, chacun avec son poids.
    # / The receipt: two items, each with its weight.
    donnees_du_ticket = formatter_ticket_vente(vente, None)
    assert len(donnees_du_ticket["articles"]) == 2
    texte_du_ticket = texte_imprime_par_l_imprimante_escpos(donnees_du_ticket)
    assert texte_du_ticket.count("0,350 kg") == 2, texte_du_ticket
    assert "0,700" not in texte_du_ticket, texte_du_ticket
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 14 — Admin : la quantité d'une pesée avec son unité
# / 14 — Admin: a weighing's quantity with its unit
# --------------------------------------------------------------------------


def test_admin_quantite_d_une_pesee_avec_l_unite(lieu):
    """
    355 g de comté à 12,90 €/kg (4,58 €) en espèces. La fiche « Vente » de l'admin
    (articles de la vente) et l'onglet des lignes (`LigneArticleInline`) montrent la
    quantité « 0,355 kg » et le prix unitaire « 12,90 €/kg » (Q-H4).
    / The admin sale page and the lines tab show "0,355 kg" and "12,90 €/kg".
    """
    comte = creer_un_comte_au_kilo()
    parcours = vendre_une_pesee(
        lieu,
        comte,
        quantite_saisie=355,
        montant_en_centimes=458,
        facon_de_payer="espece",
    )
    ligne_du_comte = seul_article_de_la_vente(parcours.vente)
    assert ligne_du_comte.qty == Decimal("0.355")

    articles_de_la_vente = ArticlesDeLaVenteInline(Vente, staff_admin_site)
    assert texte_normalise(str(articles_de_la_vente.quantite(ligne_du_comte))) == (
        "0,355 kg"
    )
    assert texte_normalise(str(articles_de_la_vente.prix_unitaire(ligne_du_comte))) == (
        "12,90 €/kg"
    )

    onglet_des_lignes = LigneArticleInline(Membership, staff_admin_site)
    assert texte_normalise(str(onglet_des_lignes.qty_decimal(ligne_du_comte))) == (
        "0,355 kg"
    )
    assert texte_normalise(str(onglet_des_lignes.amount_decimal(ligne_du_comte))) == (
        "12,90 €/kg"
    )
    verifier_egalites(parcours.vente)


# --------------------------------------------------------------------------
# 15 — Export tableur du Z : la quantité d'une pesée porte son unité
# / 15 — Z spreadsheet export: a weighing's quantity carries its unit
# --------------------------------------------------------------------------


def test_excel_quantite_d_une_pesee_avec_l_unite(lieu):
    """
    350 g de comté (stock en grammes) en espèces. L'export tableur d'une clôture dont le
    rapport contient cette seule vente : dans « Ventes par produit », la cellule de
    quantité du comté porte l'unité « kg » (dans son texte ou dans son format de
    nombre). La clôture est en mémoire, jamais enregistrée.
    / The spreadsheet export: the comté quantity cell carries the "kg" unit (text or
    number format). In-memory closure.
    """
    comte = creer_un_comte_au_kilo(avec_stock_en_grammes=True)
    parcours = vendre_une_pesee(
        lieu,
        comte,
        quantite_saisie=350,
        montant_en_centimes=452,
        facon_de_payer="espece",
    )
    rapport = rapport_de_l_instant_de_la_vente(parcours.vente)
    cloture_en_memoire = ClotureCaisseUnique(
        numero_sequentiel=1,
        niveau=ClotureCaisseUnique.NIVEAU_JOURNALIER,
        hmac_hash="empreinte-de-test",
        datetime_fin=timezone.now(),
        rapport_json={
            "chiffre_affaires": rapport.section_chiffre_affaires(),
            "detail": rapport.section_detail(),
        },
    )

    contenu_en_octets, _nom_du_fichier, _type = generer_excel_cloture(
        cloture_en_memoire
    )
    classeur = load_workbook(io.BytesIO(contenu_en_octets))
    feuille = classeur.active

    cellules_de_quantite_du_comte = []
    for rangee in feuille.iter_rows():
        for position, case in enumerate(rangee):
            case_du_nom_du_comte = case.value == comte.produit.name
            if case_du_nom_du_comte and position + 1 < len(rangee):
                cellules_de_quantite_du_comte.append(rangee[position + 1])
    assert len(cellules_de_quantite_du_comte) == 1
    case_de_la_quantite = cellules_de_quantite_du_comte[0]
    unite_dans_le_texte = (
        isinstance(case_de_la_quantite.value, str) and "kg" in case_de_la_quantite.value
    )
    unite_dans_le_format = "kg" in (case_de_la_quantite.number_format or "")
    assert unite_dans_le_texte or unite_dans_le_format, (
        f"Quantité {case_de_la_quantite.value!r}, format "
        f"{case_de_la_quantite.number_format!r} : l'unité « kg » manque."
    )
    verifier_egalites(parcours.vente)


# --------------------------------------------------------------------------
# 16 — Rapport : un produit vendu au poids et à la pièce donne deux lignes
# / 16 — Report: a product sold by weight and by piece gives two rows
# --------------------------------------------------------------------------


def test_rapport_produit_au_poids_et_a_la_piece_deux_lignes(lieu):
    """
    Le comté a deux tarifs : au kilo (12,90 €/kg) et « Portion » à la pièce (4,00 €).
    Une vente en espèces : 350 g au kilo + 1 portion. Le détail des ventes du rapport
    ne mélange jamais deux unités sur une ligne : le comté a DEUX lignes, une de
    0,350 « kg », une de 1 pièce (unité vide).
    / A product with a per-kg price and a per-piece price: two report rows, one in kg,
    one in pieces.
    """
    comte = creer_un_comte_au_kilo(avec_stock_en_grammes=True)
    tarif_portion = Price.objects.create(
        product=comte.produit,
        name="Portion",
        prix=Decimal("4.00"),
        publish=True,
    )
    point_de_vente = creer_un_point_de_vente([comte.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()
    champs_du_panier = champs_d_une_pesee(
        comte, quantite_saisie=350, montant_en_centimes=452
    )
    champs_du_panier[f"repid-{comte.produit.uuid}--{tarif_portion.uuid}"] = "1"

    reponse = payer_un_panier_compose_a_la_caisse(
        client_du_caissier,
        point_de_vente,
        champs_du_panier,
        moyen_de_paiement="espece",
        cle_d_idempotence=cle_d_idempotence,
    )

    assert reponse.status_code == 200, reponse.content.decode()[:400]
    vente = retrouver_la_vente_de_la_cle(cle_d_idempotence)
    assert vente.articles.count() == 2
    detail = rapport_de_l_instant_de_la_vente(vente).section_detail()

    lignes_du_comte = []
    for ligne_du_detail in detail["ventes_par_produit"].values():
        if ligne_du_detail["nom"] == comte.produit.name:
            lignes_du_comte.append(
                (Decimal(ligne_du_detail["quantite"]), ligne_du_detail.get("unite", ""))
            )
    assert sorted(lignes_du_comte) == sorted(
        [(Decimal("0.350"), "kg"), (Decimal("1"), "")]
    ), lignes_du_comte
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 17 — Vidage : l'article « Jetons cadeau repris » n'a pas de moyen
# / 17 — Card emptying: the "gift tokens taken back" item has no method
# --------------------------------------------------------------------------


def test_jetons_repris_au_vidage_sans_moyen_sur_la_ligne(lieu):
    """
    Le lieu n'est pas relié à l'ancien Fedow. La carte porte 5,00 € de monnaie locale
    et 2,00 € de jetons cadeau du lieu ; le caissier la vide. La vente `VIDAGE_CARTE` a
    un article « Jetons cadeau repris au vidage » (200) : sa ligne n'a pas de moyen
    (Q-H2), le règlement « jetons » (LG) le porte.
    / The "gift tokens taken back" line has no payment method; the LG payment has it.
    """
    caisse = preparer_la_caisse(lieu)
    carte_du_client = creer_une_carte_client_sans_solde(lieu)
    poser_un_solde_sur_la_carte(
        carte_du_client, monnaie_locale_propre_au_lieu(lieu), 500
    )
    poser_un_solde_sur_la_carte(
        carte_du_client, monnaie_cadeau_propre_au_lieu(lieu), 200
    )

    reponse, _classe_fedow_api = vider_la_carte_a_la_caisse(
        caisse, carte_du_client, MagicMock(), lieu_relie=False
    )

    assert reponse.status_code == 200
    vente = retrouver_la_vente_de_vidage(carte_du_client)
    lignes_de_la_vente = list(vente.articles.all())
    assert len(lignes_de_la_vente) == 1
    ligne_des_jetons_repris = lignes_de_la_vente[0]
    assert ligne_des_jetons_repris.total_catalogue == 200
    assert champ_vide(ligne_des_jetons_repris.payment_method), (
        f"Moyen de la ligne : {ligne_des_jetons_repris.payment_method!r}."
    )
    assert vente.reglements.filter(moyen=PaymentMethod.LOCAL_GIFT).count() == 1
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 18 — Garde de la cascade : un débit faux n'écrit rien
# / 18 — Cascade guard: a wrong debit writes nothing
# --------------------------------------------------------------------------


def test_garde_de_la_cascade_debit_faux_rien_d_ecrit(lieu):
    """
    Un jus à 3,50 € reçoit un débit de 3,00 € seulement : la somme des débits de
    l'article ne vaut pas son total catalogue (350). `_creer_lignes_articles_cascade`
    lève `EgaliteDeVenteRompue` et n'écrit aucune ligne.
    / A 3.50 € juice with a 3.00 € debit: EgaliteDeVenteRompue, no line written.
    """
    jus = creer_un_article_de_caisse("jus", prix_en_euros="3.50", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([jus.produit])
    donnees_du_panier = QueryDict(mutable=True)
    donnees_du_panier[f"repid-{cle_de_panier(jus)}"] = "1"
    articles_du_panier = _extraire_articles_du_panier(donnees_du_panier, point_de_vente)
    assert len(articles_du_panier) == 1
    article_du_jus = articles_du_panier[0]
    vente = ouvrir_vente(origine=SaleOrigin.LABOUTIK, nature=Vente.Nature.VENTE)
    debit_faux = (article_du_jus, None, 300, PaymentMethod.CASH)

    with pytest.raises(EgaliteDeVenteRompue):
        _creer_lignes_articles_cascade(
            lignes_pre_calculees=[debit_faux],
            vente=vente,
            uuid_transaction=uuid_module.uuid4(),
            point_de_vente=point_de_vente,
        )

    assert vente.articles.count() == 0
    assert not LigneArticle.objects.filter(pricesold__price=jus.tarif).exists()


# --------------------------------------------------------------------------
# 19 — Retour de consigne remboursé sur la carte : colonnes vides
# / 19 — Deposit return refunded on the card: empty columns
# --------------------------------------------------------------------------


def test_retour_consigne_nfc_sans_moyen_ni_carte(lieu):
    """
    Un gobelet rendu, remboursé sur la carte du client (`_rembourser_consigne_par_nfc`) :
    la ligne n'a ni moyen, ni monnaie, ni carte, ni portefeuille (Q-H2) ; le règlement
    négatif en monnaie locale porte la carte.
    / A cup refunded on the card: the line has no method, currency, card nor wallet.
    """
    consigne = creer_un_gobelet_et_son_retour(lieu)
    point_de_vente = creer_un_point_de_vente([consigne.retour.produit])
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=0)
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()

    reponse = payer_a_la_caisse(
        client_du_caissier,
        point_de_vente,
        consigne.retour,
        quantite=1,
        moyen_de_paiement="nfc",
        autres_champs={
            "tag_id": carte_du_client.tag_id,
            "cle_idempotence_paiement": cle_d_idempotence,
        },
    )

    assert reponse.status_code == 200
    vente = retrouver_la_vente_de_la_cle(cle_d_idempotence)
    ligne_du_retour = seul_article_de_la_vente(vente)
    assert champ_vide(ligne_du_retour.payment_method)
    assert champ_vide(ligne_du_retour.asset)
    assert ligne_du_retour.carte_id is None
    assert ligne_du_retour.wallet_id is None
    reglement_du_retour = vente.reglements.get()
    assert reglement_du_retour.carte_id == carte_du_client.pk
    verifier_egalites(vente)


# --------------------------------------------------------------------------
# 20 — Cascade : le stock est décrémenté du poids saisi
# / 20 — Cascade: stock decremented by the entered weight
# --------------------------------------------------------------------------


def test_stock_decremente_par_le_poids_en_cascade(lieu):
    """
    350 g de comté (stock de 10 000 g) payés par la carte : 2,00 € de jetons cadeau
    puis la monnaie locale. Le stock perd 350 g (le poids saisi, `weight_quantity`),
    une fois.
    / A weighing paid by card in two currencies: the stock loses 350 g, once.
    """
    comte = creer_un_comte_au_kilo(avec_stock_en_grammes=True)

    parcours = vendre_une_pesee(
        lieu,
        comte,
        quantite_saisie=350,
        montant_en_centimes=452,
        facon_de_payer="nfc_deux_monnaies",
    )

    stock_du_comte = Stock.objects.get(product=comte.produit)
    assert stock_du_comte.quantite == 10000 - 350
    verifier_egalites(parcours.vente)


# --------------------------------------------------------------------------
# 21 — Commande de table : un tarif au poids (ou un prix libre) est refusé
# / 21 — Table order: a weight price (or a free price) is refused
# --------------------------------------------------------------------------


@pytest.mark.parametrize("sorte_de_tarif", ["au_poids", "prix_libre"])
def test_commande_de_table_refuse_un_tarif_au_poids(lieu, sorte_de_tarif):
    """
    Ouvrir une commande de table avec un tarif au poids (ou un prix libre, sans
    montant) : la commande ne connaît ni le poids ni le montant. Refus (400), aucune
    commande ni article n'est écrit.
    / Opening a table order with a weight price (or a free price): refused (400),
    nothing written.
    """
    if sorte_de_tarif == "au_poids":
        article = creer_un_comte_au_kilo()
    else:
        article = creer_un_article_de_caisse(
            "don libre", prix_en_euros="1.00", taux_tva="0.00", prix_libre=True
        )
    point_de_vente = creer_un_point_de_vente([article.produit])
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)

    reponse = client_du_caissier.post(
        "/laboutik/commande/ouvrir/",
        data={
            "uuid_pv": str(point_de_vente.uuid),
            "articles": [
                {
                    "product_uuid": str(article.produit.uuid),
                    "price_uuid": str(article.tarif.uuid),
                    "qty": 1,
                }
            ],
        },
        content_type="application/json",
    )

    assert reponse.status_code == 400, reponse.content.decode()[:400]
    assert not ArticleCommandeSauvegarde.objects.filter(price=article.tarif).exists()


@pytest.mark.parametrize("sorte_de_tarif", ["au_poids", "prix_libre"])
def test_ajout_a_une_commande_ouverte_refuse_un_tarif_au_poids(lieu, sorte_de_tarif):
    """
    Une table a une commande ouverte avec un jus (article ordinaire). Y ajouter un
    tarif au poids (ou un prix libre, sans montant) est refusé (400) : la commande ne
    connaît ni le poids ni le montant. Rien n'est ajouté, le jus reste seul.
    / Adding a weight (or free) price to an open order: refused (400), nothing added.
    """
    jus = creer_un_article_de_caisse("jus", prix_en_euros="3.50", taux_tva="20.00")
    table_et_commande = creer_une_table_occupee_avec_sa_commande(jus, quantite=1)
    if sorte_de_tarif == "au_poids":
        article = creer_un_comte_au_kilo()
    else:
        article = creer_un_article_de_caisse(
            "don libre", prix_en_euros="1.00", taux_tva="0.00", prix_libre=True
        )
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)

    reponse = client_du_caissier.post(
        f"/laboutik/commande/ajouter/{table_et_commande.commande.uuid}/",
        data=[
            {
                "product_uuid": str(article.produit.uuid),
                "price_uuid": str(article.tarif.uuid),
                "qty": 1,
            }
        ],
        content_type="application/json",
    )

    assert reponse.status_code == 400, reponse.content.decode()[:400]
    assert not ArticleCommandeSauvegarde.objects.filter(price=article.tarif).exists()
    assert table_et_commande.commande.articles.count() == 1


# --------------------------------------------------------------------------
# 22 — Imprimante Sunmi intégrée : le prix au kilo d'une pesée
# / 22 — Built-in Sunmi printer: a weighing's price per kilo
# --------------------------------------------------------------------------


def test_sunmi_interne_imprime_le_prix_au_kilo(lieu):
    """
    350 g de comté à 12,90 €/kg en espèces : le ticket de l'imprimante Sunmi intégrée
    imprime le détail du poids, « 0,350 kg » et « 12,90 €/kg », comme l'imprimante
    ESC/POS.
    / The built-in Sunmi receipt prints "0,350 kg" and "12,90 €/kg".
    """
    comte = creer_un_comte_au_kilo(avec_stock_en_grammes=True)
    parcours = vendre_une_pesee(
        lieu,
        comte,
        quantite_saisie=350,
        montant_en_centimes=452,
        facon_de_payer="espece",
    )

    donnees_du_ticket = formatter_ticket_vente(parcours.vente, None)
    texte_du_ticket = texte_imprime_par_l_imprimante_sunmi(donnees_du_ticket)

    assert "0,350 kg" in texte_du_ticket, texte_du_ticket
    assert "12,90" in texte_du_ticket and "/kg" in texte_du_ticket, texte_du_ticket
    verifier_egalites(parcours.vente)


# --------------------------------------------------------------------------
# 23 — Outil de test reset_carte : une carte qui a des règlements
# / 23 — Test tool reset_carte: a card with payments
# --------------------------------------------------------------------------


def test_reset_carte_ne_plante_plus(lieu):
    """
    Une carte NFC paie une bière par la caisse : sa vente a un règlement qui porte le
    portefeuille de la carte (`Reglement.wallet`, protégé). `reset_carte` (outil de
    test, DEBUG seulement) remet la carte à zéro sans erreur : plus d'utilisateur ni
    de portefeuille éphémère sur la carte. La vente et son règlement restent.
    / A card with a sale payment: reset_carte runs without error and leaves the card
    anonymous; the sale and its payment stay.
    """
    biere = creer_un_article_de_caisse("biere", prix_en_euros="5.00", taux_tva="20.00")
    point_de_vente = creer_un_point_de_vente([biere.produit])
    carte_du_client = creer_une_carte_nfc_chargee(lieu, solde_en_centimes=1000)
    client_du_caissier = creer_un_administrateur_du_lieu(lieu)
    cle_d_idempotence = nouvelle_cle_d_idempotence()
    reponse = payer_par_la_carte_du_client(
        client_du_caissier,
        point_de_vente,
        {f"repid-{cle_de_panier(biere)}": "1"},
        carte_du_client,
        cle_d_idempotence,
    )
    assert reponse.status_code == 200
    vente = retrouver_la_vente_de_la_cle(cle_d_idempotence)
    assert vente.reglements.filter(wallet__isnull=False).exists()

    with override_settings(DEBUG=True):
        resultat_du_reset = reset_carte(tag_id=carte_du_client.tag_id)

    assert resultat_du_reset["status"] == "ok"
    carte_du_client.refresh_from_db()
    assert carte_du_client.user_id is None
    assert carte_du_client.wallet_ephemere_id is None
    assert Vente.objects.filter(pk=vente.pk).exists()
    assert vente.reglements.count() == 1
