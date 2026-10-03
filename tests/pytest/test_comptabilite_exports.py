"""
tests/pytest/test_comptabilite_exports.py
Les exports d'une clôture (PDF, tableur Excel, CSV) et son email : ils lisent le rapport
stocké (`rapport_json`) et montrent les mêmes sections que la fiche de l'admin.
/ The exports of a closure (PDF, Excel, CSV) and its email read the stored report and
show the same sections as the admin page.

LOCALISATION : tests/pytest/test_comptabilite_exports.py

RÈGLES MÉTIER TESTÉES (fiche F §2)
- Chaque export contient toutes les sections du rapport stocké, l'essentiel d'abord
  (en-tête, chiffre d'affaires, règlements, caisse espèces, réconciliation, offerts),
  puis l'annexe, les points, la marge brute, le détail, l'intégrité.
- Les totaux imprimés sont ceux du rapport stocké, au centime, sur un scénario mêlé :
  espèces, CB, monnaie locale, un avoir, une carte vidée.
- Le tableur écrit les montants en NOMBRES (en euros) : on peut les additionner.
- L'email de clôture montre le chiffre d'affaires TTC, avec son vrai libellé.
- Le FEC a la forme d'un FEC et ses écritures sont équilibrées (le détail de la
  ventilation : tests/pytest/test_fec_equilibre.py). Le FEC est le seul export
  comptable.
/ Business rules tested: every export has every section, totals equal the stored
report on a mixed scenario, numbers in the spreadsheet, the email shows the revenue.

SCHÉMA DÉDIÉ
Une clôture lit TOUTES les ventes du lieu : il faut un lieu qui ne contient que les
ventes du test (`FastTenantTestCase`). Chaque test annule sa transaction à la fin. Le
singleton `LaboutikConfiguration` est créé en base (tests/PIEGES.md 9.86). Les ventes
sont écrites PAR LE SERVICE (`fabriques_vente.py`), puis la J glissante est créée par la
tâche de clôture (heure remplacée pendant l'appel, comme test_cloture_unique.py).
/ Dedicated schema; real sales through the service, then the sliding J.

Spécification : TECH_DOC/SESSIONS/COMPTABILITE/CHANTIER-05-F-rapport-unique.md (§2).

Lancer / Run : make test ARGS="tests/pytest/test_comptabilite_exports.py"
"""

import sys

# Le code Django vit dans /DjangoFiles a l'interieur du conteneur.
# / Django code lives in /DjangoFiles inside the container.
sys.path.insert(0, "/DjangoFiles")

import django  # noqa: E402

django.setup()

import csv  # noqa: E402
import io  # noqa: E402
import re  # noqa: E402
import uuid  # noqa: E402
from datetime import datetime, time, timedelta  # noqa: E402
from decimal import Decimal  # noqa: E402
from unittest.mock import MagicMock, patch  # noqa: E402
from zoneinfo import ZoneInfo  # noqa: E402

from django.contrib.messages.storage.fallback import FallbackStorage  # noqa: E402
from django.contrib.sessions.middleware import SessionMiddleware  # noqa: E402
from django.db import connection  # noqa: E402
from django.template.loader import render_to_string  # noqa: E402
from django.test import RequestFactory  # noqa: E402
from django.utils import timezone, translation  # noqa: E402
from django_tenants.test.cases import FastTenantTestCase  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import comptabilite.tasks  # noqa: E402
from Administration.admin.site import staff_admin_site  # noqa: E402
from AuthBillet.models import TibilletUser, Wallet  # noqa: E402
from BaseBillet.models import (  # noqa: E402
    CategorieProduct,
    Configuration,
    Event,
    LigneArticle,
    PaymentMethod,
    PostalAddress,
    Price,
    PriceSold,
    Product,
    ProductSold,
    Reservation,
    SaleOrigin,
    Tva,
)
from BaseBillet.models_vente import Vente  # noqa: E402
from BaseBillet.services_vente import ecrire_la_vente_d_avoir_d_une_ligne  # noqa: E402
from comptabilite.admin import ClotureCaisseAdmin  # noqa: E402
from comptabilite.csv_export import generer_csv_cloture  # noqa: E402
from comptabilite.excel_export import generer_excel_cloture  # noqa: E402
from comptabilite.integrite import verifier_chaine_clotures  # noqa: E402
from comptabilite.models import ClotureCaisse  # noqa: E402
from comptabilite.pdf import generer_pdf_cloture, html_du_pdf_de_la_cloture  # noqa: E402
from comptabilite.presentation import sections_pour_affichage  # noqa: E402
from fabriques_panier import (  # noqa: E402
    creer_utilisateur,
    identifiant_unique,
    taches_celery_enregistrees,
)
from fabriques_vente import (  # noqa: E402
    creer_tarif_vendu,
    fabriquer_vente_encaissee,
    verifier_egalites,
)
from fedow_core.models import Asset  # noqa: E402
from laboutik.models import (  # noqa: E402
    CompteComptable,
    LaboutikConfiguration,
    PointDeVente,
)

# L'espace insécable : entre les milliers, et entre le nombre et « € ».
# / The non-breaking space.
ESPACE_INSECABLE = "\u00a0"

# Les titres des sections d'une J, dans l'ordre de la fiche F §2.
# / The section titles of a J, in the order of sheet F §2.
TITRES_DES_SECTIONS_D_UNE_J = [
    "En-tête",
    "Chiffre d'affaires",
    "Règlements",
    "Caisse espèces",
    "Réconciliation",
    "Offerts",
    "Annexe : avoirs, recharges, écarts, corrections",
    "Points",
    "Marge brute",
    "Détail des ventes",
    "Intégrité",
]


def euros(nombre_en_texte):
    """
    Un montant attendu, écrit à la main : « 17,50 » → « 17,50 € » avec l'espace
    insécable. Le texte est écrit par le test, jamais par le code testé.
    / An expected amount, hand-written.
    """
    nombre_avec_espaces_insecables = nombre_en_texte.replace(" ", ESPACE_INSECABLE)
    return f"{nombre_avec_espaces_insecables}{ESPACE_INSECABLE}€"


class TestComptabiliteExports(FastTenantTestCase):
    """
    Les exports PDF, Excel et CSV d'une J d'un scénario mêlé, et l'email de clôture.
    / PDF, Excel and CSV exports of a mixed-scenario J, and the closure email.
    """

    @classmethod
    def get_test_schema_name(cls):
        return "test_comptabilite_exports"

    @classmethod
    def get_test_tenant_domain(cls):
        return "test-comptabilite-exports.tibillet.localhost"

    @classmethod
    def setup_tenant(cls, tenant):
        """`Client.name` est unique et obligatoire. / Client.name is unique."""
        tenant.name = "Test comptabilite exports"

    def setUp(self):
        """
        Le lieu de test (Paris, fermeture à 2 h, aucun email de rapport), puis le
        scénario mêlé et sa J :
        - trois jus (10,50 €) payés 5,00 € en monnaie locale et 5,50 € en CB ;
        - deux jus en espèces (7,00 €) ;
        - un jus en espèces (3,50 €), remboursé en espèces par un avoir (−3,50 €) ;
        - une carte de 2,00 € vidée, les espèces rendues (−2,00 €).
        Chiffre d'affaires TTC : 1050 + 700 + 350 − 350 = 1750.
        / The test venue, then the mixed scenario and its J. Revenue 1750.
        """
        # Le rollback du test précédent a rendu le `search_path` au public.
        # / The previous test's rollback returned the search_path to public.
        connection.set_tenant(self.tenant)
        LaboutikConfiguration.get_solo().save()

        # Les mentions légales du lieu, figées dans le Z à la clôture. Toutes
        # les valeurs lues sont écrites ici : le cache de la configuration survit à
        # l'annulation du test précédent.
        # / The venue's legal mentions, stored in the Z; every value is written here.
        with taches_celery_enregistrees():
            adresse_du_lieu = PostalAddress.objects.create(
                name="Adresse exports",
                street_address="7 place du Marché",
                address_locality="Lyon",
                postal_code="69001",
                address_country="FR",
            )
        configuration = Configuration.get_solo()
        configuration.fuseau_horaire = "Europe/Paris"
        configuration.heure_de_fermeture = time(2, 0)
        configuration.rapport_emails = ""
        configuration.postal_address = adresse_du_lieu
        configuration.postal_code = 69001
        configuration.city = "Lyon"
        configuration.siren = "552100554"
        configuration.tva_number = "FR40552100554"
        configuration.phone = "0400000000"
        configuration.save()

        self.tarif_du_jus = creer_tarif_vendu(
            nom="Jus",
            prix_en_euros="3.50",
            taux_tva="20.00",
            methode_caisse=Product.VENTE,
        )
        self.point_de_vente = PointDeVente.objects.create(
            name="Bar", code_journal="", hidden=True
        )
        self.monnaie_locale = self._monnaie_locale()

        self._vendre_trois_jus_en_monnaie_locale_et_cb()
        self._vendre_des_jus_en_especes(quantite=2)
        vente_remboursee = self._vendre_des_jus_en_especes(quantite=1)
        self._rembourser_en_especes(vente_remboursee)
        self._vider_une_carte(200)

        self.cloture = self._cloturer_la_journee()
        self.rapport = self.cloture.rapport_json

    # ------------------------------------------------------------------
    # Outils du test
    # / Test helpers
    # ------------------------------------------------------------------

    def _monnaie_locale(self):
        """Une monnaie locale créée par le lieu de test. / A local currency."""
        portefeuille_d_origine = Wallet.objects.create(
            name=f"Portefeuille exports {identifiant_unique()}"
        )
        return Asset.objects.create(
            name="Monnaie locale exports",
            currency_code="EUR",
            category=Asset.TLF,
            tenant_origin=self.tenant,
            wallet_origin=portefeuille_d_origine,
        )

    def _vendre_trois_jus_en_monnaie_locale_et_cb(self):
        """
        Trois jus à 3,50 €, payés 5,00 € en monnaie locale et 5,50 € en CB, en deux
        « parts » comme la caisse les écrit pendant la transition.
        / Three juices, 5.00 € local currency + 5.50 € card, in two parts.
        """
        vente = fabriquer_vente_encaissee(
            origine=SaleOrigin.LABOUTIK,
            point_de_vente=self.point_de_vente,
            articles=[
                {
                    "pricesold": self.tarif_du_jus,
                    "quantite": Decimal("1.428571"),
                    "prix_unitaire": 350,
                    "taux_tva": Decimal("20"),
                    "total_catalogue_impose": 500,
                },
                {
                    "pricesold": self.tarif_du_jus,
                    "quantite": Decimal("1.571429"),
                    "prix_unitaire": 350,
                    "taux_tva": Decimal("20"),
                    "total_catalogue_impose": 550,
                },
            ],
            reglements=[
                {
                    "moyen": PaymentMethod.LOCAL_EURO,
                    "montant": 500,
                    "asset": self.monnaie_locale.uuid,
                },
                {"moyen": PaymentMethod.CC, "montant": 550},
            ],
        )
        verifier_egalites(vente)

    def _vendre_des_jus_en_especes(self, quantite):
        """
        `quantite` jus à 3,50 € en espèces, sur le point de vente ; la ligne est VALID.
        / Juices in cash at the point of sale; the line is VALID.
        """
        vente = fabriquer_vente_encaissee(
            origine=SaleOrigin.LABOUTIK,
            point_de_vente=self.point_de_vente,
            articles=[
                {
                    "pricesold": self.tarif_du_jus,
                    "quantite": Decimal(quantite),
                    "prix_unitaire": 350,
                    "taux_tva": Decimal("20"),
                    "status": LigneArticle.VALID,
                },
            ],
            reglements=[{"moyen": PaymentMethod.CASH, "montant": 350 * quantite}],
        )
        verifier_egalites(vente)
        return vente

    def _rembourser_en_especes(self, vente):
        """L'avoir de la ligne, remboursé en espèces. / Cash credit note of the line."""
        ligne_vendue = LigneArticle.objects.get(vente=vente)
        with taches_celery_enregistrees():
            article_d_avoir = ecrire_la_vente_d_avoir_d_une_ligne(
                ligne_vendue,
                quantite=Decimal("1"),
                moyen_rembourse=PaymentMethod.CASH,
                origine=SaleOrigin.ADMIN,
            )
        verifier_egalites(Vente.objects.get(pk=article_d_avoir.vente_id))

    def _vider_une_carte(self, montant_en_centimes):
        """Une carte vidée, les espèces rendues. / A card emptied, cash given back."""
        vente_du_vidage = fabriquer_vente_encaissee(
            origine=SaleOrigin.LABOUTIK,
            nature=Vente.Nature.VIDAGE_CARTE,
            point_de_vente=self.point_de_vente,
            reglements=[
                {
                    "moyen": PaymentMethod.LOCAL_EURO,
                    "montant": montant_en_centimes,
                    "asset": self.monnaie_locale.uuid,
                },
                {"moyen": PaymentMethod.CASH, "montant": -montant_en_centimes},
            ],
        )
        verifier_egalites(vente_du_vidage)

    def _cloturer_la_journee(self):
        """Le « Z de fin de service », une minute après les ventes. / The J."""
        moment_de_la_cloture = timezone.now() + timedelta(minutes=1)
        with patch("django.utils.timezone.now", return_value=moment_de_la_cloture):
            uuid_de_la_cloture = comptabilite.tasks.generer_cloture_pour_tenant(
                schema_name=self.tenant.schema_name,
                niveau=ClotureCaisse.NIVEAU_JOURNALIER,
            )
        return ClotureCaisse.objects.get(uuid=uuid_de_la_cloture)

    def _lignes_du_csv(self, cloture=None):
        """
        Les lignes du CSV d'une clôture (celle du scénario par défaut), en français.
        / The CSV rows of a closure, in French.
        """
        if cloture is None:
            cloture = self.cloture
        with translation.override("fr"):
            contenu_en_octets, _nom, _type = generer_csv_cloture(cloture)
        contenu = contenu_en_octets.decode("utf-8-sig")
        return list(csv.reader(io.StringIO(contenu), delimiter=";"))

    def _ligne_qui_commence_par(self, lignes, libelle):
        """La première ligne dont la 1ʳᵉ cellule vaut `libelle`. / First row by label."""
        for ligne in lignes:
            if len(ligne) > 0 and ligne[0] == libelle:
                return ligne
        raise AssertionError(f"Ligne absente : {libelle}")

    def _feuille_du_tableur(self, cloture=None):
        """
        La feuille du tableur d'une clôture (celle du scénario par défaut), en français.
        / The spreadsheet sheet of a closure, in French.
        """
        if cloture is None:
            cloture = self.cloture
        with translation.override("fr"):
            contenu_en_octets, _nom, _type = generer_excel_cloture(cloture)
        classeur = load_workbook(io.BytesIO(contenu_en_octets))
        return classeur.active

    def _valeurs_de_la_ligne_du_tableur(self, feuille, libelle):
        """
        Les valeurs de la première ligne du tableur dont la colonne A vaut `libelle`.
        / The values of the first sheet row whose column A is `libelle`.
        """
        for ligne in feuille.iter_rows(values_only=True):
            if ligne[0] == libelle:
                return ligne
        raise AssertionError(f"Ligne absente du tableur : {libelle}")

    def _html_du_pdf(self, cloture=None):
        """Le HTML que WeasyPrint imprime, en français. / The HTML printed to PDF."""
        if cloture is None:
            cloture = self.cloture
        with translation.override("fr"):
            return html_du_pdf_de_la_cloture(cloture)

    def _page_de_la_fiche(self, cloture):
        """
        Le HTML des sections de la fiche admin d'une clôture, en français.
        / The admin page sections of a closure, in French.
        """
        with translation.override("fr"):
            sections = sections_pour_affichage(cloture.rapport_json)
            return render_to_string(
                "comptabilite/admin/_sections_rapport.html", {"sections": sections}
            )

    def _tarif_vendu_d_un_produit(self, nom, uuid_du_produit, categorie_article):
        """
        Un produit au nom EXACT (sans préfixe de test) et à l'uuid choisi, son tarif et
        son tarif vendu.
        / A product with an EXACT name and a chosen uuid, its price and sold price.
        """
        tva_a_vingt_pour_cent, _tva_creee = Tva.objects.get_or_create(
            tva_rate=Decimal("20.00")
        )
        produit = Product.objects.create(
            uuid=uuid_du_produit,
            name=nom,
            categorie_article=categorie_article,
            tva=tva_a_vingt_pour_cent,
        )
        tarif = Price.objects.create(
            product=produit, name="Tarif unique", prix=Decimal("3.50"), publish=True
        )
        produit_vendu = ProductSold.objects.create(product=produit)
        return PriceSold.objects.create(
            productsold=produit_vendu, price=tarif, prix=tarif.prix
        )

    def _vendre_en_especes_a(self, moment, tarif_vendu, reservation=None):
        """
        Une vente d'un article à 3,50 € en espèces, encaissée à `moment`.
        / A 3.50 € cash sale settled at `moment`.
        """
        article = {
            "pricesold": tarif_vendu,
            "quantite": Decimal("1"),
            "prix_unitaire": 350,
            "taux_tva": Decimal("20"),
        }
        if reservation is not None:
            article["reservation"] = reservation
        with patch("django.utils.timezone.now", return_value=moment):
            vente = fabriquer_vente_encaissee(
                origine=SaleOrigin.LABOUTIK,
                articles=[article],
                reglements=[{"moyen": PaymentMethod.CASH, "montant": 350}],
            )
        verifier_egalites(vente)

    def _cloturer_a(self, moment):
        """
        Une J créée à `moment` (après la J du scénario), relue depuis la base.
        / A J created at `moment`, read back from the database.
        """
        with patch("django.utils.timezone.now", return_value=moment):
            uuid_de_la_cloture = comptabilite.tasks.generer_cloture_pour_tenant(
                schema_name=self.tenant.schema_name,
                niveau=ClotureCaisse.NIVEAU_JOURNALIER,
            )
        return ClotureCaisse.objects.get(uuid=uuid_de_la_cloture)

    def _billet_d_un_evenement_vendu_a(self, moment, nom, date, uuid_de_l_evenement):
        """
        Un événement à l'uuid choisi, et un billet vendu pour lui (réservation valide),
        encaissé à `moment`.
        / An event with a chosen uuid, and one ticket sold for it at `moment`.
        """
        with taches_celery_enregistrees():
            evenement = Event.objects.create(
                uuid=uuid_de_l_evenement,
                name=nom,
                datetime=date,
                end_datetime=date + timedelta(hours=2),
                jauge_max=100,
            )
        tarif_vendu_du_billet = self._tarif_vendu_d_un_produit(
            f"Entrée {identifiant_unique()}", uuid.uuid4(), Product.BILLET
        )
        tarif_vendu_du_billet.productsold.event = evenement
        tarif_vendu_du_billet.productsold.save()
        reservation = Reservation.objects.create(
            user_commande=creer_utilisateur(),
            event=evenement,
            status=Reservation.VALID,
        )
        self._vendre_en_especes_a(moment, tarif_vendu_du_billet, reservation)

    def _position_de_la_ligne_qui_contient(self, lignes, texte):
        """
        La position de la première ligne (CSV ou tableur) dont une cellule vaut `texte`.
        / Position of the first row with a cell equal to `texte`.
        """
        for position, ligne in enumerate(lignes):
            if texte in ligne:
                return position
        raise AssertionError(f"Texte absent : {texte}")

    # ------------------------------------------------------------------
    # Le scénario
    # / The scenario
    # ------------------------------------------------------------------

    def test_le_scenario_mele_donne_le_rapport_attendu(self):
        """
        Le rapport stocké du scénario : CA 1750 ; argent (section 3) CB 550 + espèces
        700 + 350 − 350 = 1250 ; monnaie locale 500 ; argent reçu 1250 − 200 (carte
        vidée) = 1050 ; un avoir de 350 ; une carte vidée.
        / The stored report of the scenario.
        """
        assert self.rapport["chiffre_affaires"]["total_ttc_en_centimes"] == 1750
        assert self.rapport["reglements"]["argent"]["total_en_centimes"] == 1250
        assert self.rapport["reglements"]["cashless"]["total_en_centimes"] == 500
        assert self.rapport["reconciliation"]["argent_recu_en_centimes"] == 1050
        assert self.rapport["annexe"]["avoirs"]["total_en_centimes"] == -350
        assert self.rapport["annexe"]["recharges_et_cartes"]["cartes_videes"]["nombre"] == 1

    # ------------------------------------------------------------------
    # CSV
    # / CSV
    # ------------------------------------------------------------------

    def test_csv_contient_toutes_les_sections_dans_l_ordre(self):
        """
        Le CSV a un titre « [Section] » par section du rapport, dans l'ordre de la
        fiche.
        / The CSV has one "[Section]" title per section, in the sheet's order.
        """
        lignes = self._lignes_du_csv()
        premieres_cellules = []
        for ligne in lignes:
            if len(ligne) > 0:
                premieres_cellules.append(ligne[0])

        positions_des_titres = []
        for titre in TITRES_DES_SECTIONS_D_UNE_J:
            titre_entre_crochets = f"[{titre}]"
            assert titre_entre_crochets in premieres_cellules, titre
            positions_des_titres.append(premieres_cellules.index(titre_entre_crochets))
        assert positions_des_titres == sorted(positions_des_titres)

    def test_csv_totaux_egaux_au_rapport(self):
        """
        Les totaux du CSV sont ceux du rapport stocké : chiffre d'affaires TTC 17,50 €,
        argent 12,50 €, cashless 5,00 €, argent reçu 10,50 €, et la phrase de
        réconciliation (sorties en positif).
        / CSV totals equal the stored report.
        """
        lignes = self._lignes_du_csv()

        ligne_du_ca = self._ligne_qui_commence_par(lignes, "Chiffre d'affaires TTC")
        assert ligne_du_ca[1] == euros("17,50")
        ligne_de_l_argent = self._ligne_qui_commence_par(lignes, "Total argent")
        assert ligne_de_l_argent[-1] == euros("12,50")
        ligne_du_cashless = self._ligne_qui_commence_par(lignes, "Total cashless")
        assert ligne_du_cashless[-1] == euros("5,00")
        ligne_de_l_argent_recu = self._ligne_qui_commence_par(lignes, "Argent reçu")
        assert ligne_de_l_argent_recu[1] == euros("10,50")

        phrase_attendue = (
            f"Argent reçu {euros('10,50')} = ventes payées en argent {euros('16,00')} "
            f"+ recharges {euros('0,00')} − remboursements {euros('3,50')} "
            f"− cartes vidées {euros('2,00')} + écarts d'encaissement {euros('0,00')}"
        )
        self._ligne_qui_commence_par(lignes, phrase_attendue)

    def test_csv_par_l_admin_est_un_fichier_joint(self):
        """
        L'admin rend le CSV en pièce jointe (`text/csv`).
        / The admin returns the CSV as an attachment.
        """
        email_de_l_administrateur = f"admin-{identifiant_unique()}@tibillet.localhost"
        administrateur = TibilletUser.objects.create(
            email=email_de_l_administrateur,
            username=email_de_l_administrateur,
            is_staff=True,
            is_superuser=True,
        )
        requete = RequestFactory().get(
            f"/admin/comptabilite/cloturecaisse/{self.cloture.uuid}/exporter-csv/"
        )
        requete.user = administrateur
        SessionMiddleware(lambda requete_recue: None).process_request(requete)
        requete._messages = FallbackStorage(requete)
        admin_des_clotures = ClotureCaisseAdmin(ClotureCaisse, staff_admin_site)

        reponse = admin_des_clotures.exporter_csv(requete, self.cloture.uuid)

        assert reponse.status_code == 200
        assert "text/csv" in reponse["Content-Type"]
        assert "attachment" in reponse["Content-Disposition"]
        assert ".csv" in reponse["Content-Disposition"]

    # ------------------------------------------------------------------
    # Tableur (Excel)
    # / Spreadsheet (Excel)
    # ------------------------------------------------------------------

    def test_tableur_contient_toutes_les_sections_dans_l_ordre(self):
        """
        Le tableur a un titre par section du rapport, dans l'ordre de la fiche.
        / The spreadsheet has one title per section, in the sheet's order.
        """
        feuille = self._feuille_du_tableur()
        premieres_cellules = []
        for ligne in feuille.iter_rows(values_only=True):
            premieres_cellules.append(ligne[0])

        positions_des_titres = []
        for titre in TITRES_DES_SECTIONS_D_UNE_J:
            assert titre in premieres_cellules, titre
            positions_des_titres.append(premieres_cellules.index(titre))
        assert positions_des_titres == sorted(positions_des_titres)

    def test_tableur_totaux_en_nombres_egaux_au_rapport(self):
        """
        Le tableur écrit les montants en nombres (euros) : chiffre d'affaires TTC 17,5 ;
        argent 12,5 ; argent reçu 10,5 ; ils valent les centimes du rapport / 100.
        / The spreadsheet writes amounts as numbers equal to the report's cents / 100.
        """
        feuille = self._feuille_du_tableur()

        ligne_du_ca = self._valeurs_de_la_ligne_du_tableur(
            feuille, "Chiffre d'affaires TTC"
        )
        ca_en_centimes = self.rapport["chiffre_affaires"]["total_ttc_en_centimes"]
        assert Decimal(str(ligne_du_ca[1])) * 100 == ca_en_centimes

        ligne_de_l_argent = self._valeurs_de_la_ligne_du_tableur(feuille, "Total argent")
        argent_en_centimes = self.rapport["reglements"]["argent"]["total_en_centimes"]
        valeurs_non_vides = []
        for valeur in ligne_de_l_argent:
            if valeur is not None:
                valeurs_non_vides.append(valeur)
        assert Decimal(str(valeurs_non_vides[-1])) * 100 == argent_en_centimes

        ligne_de_l_argent_recu = self._valeurs_de_la_ligne_du_tableur(
            feuille, "Argent reçu"
        )
        argent_recu = self.rapport["reconciliation"]["argent_recu_en_centimes"]
        assert Decimal(str(ligne_de_l_argent_recu[1])) * 100 == argent_recu

    # ------------------------------------------------------------------
    # PDF
    # / PDF
    # ------------------------------------------------------------------

    def test_pdf_contient_toutes_les_sections_dans_l_ordre(self):
        """
        Le HTML imprimé en PDF a un titre par section du rapport, dans l'ordre.
        / The printed HTML has one title per section, in order.
        """
        html = self._html_du_pdf()

        positions_des_titres = []
        for titre in TITRES_DES_SECTIONS_D_UNE_J:
            titre_echappe = titre.replace("'", "&#x27;")
            titre_html = f"<h2>{titre_echappe}</h2>"
            assert titre_html in html, titre
            positions_des_titres.append(html.index(titre_html))
        assert positions_des_titres == sorted(positions_des_titres)

    def test_pdf_totaux_egaux_au_rapport(self):
        """
        Le PDF imprime le chiffre d'affaires TTC (17,50 €), l'argent (12,50 €) et la
        phrase de réconciliation.
        / The PDF prints the revenue, the money and the reconciliation sentence.
        """
        html = self._html_du_pdf()

        assert euros("17,50") in html
        assert euros("12,50") in html
        phrase_attendue = (
            f"Argent reçu {euros('10,50')} = ventes payées en argent {euros('16,00')} "
            f"+ recharges {euros('0,00')} − remboursements {euros('3,50')} "
            f"− cartes vidées {euros('2,00')} + écarts d&#x27;encaissement "
            f"{euros('0,00')}"
        )
        assert phrase_attendue in html

    def test_pdf_est_un_vrai_pdf(self):
        """
        Le fichier rendu est un PDF (signature « %PDF- »), nommé d'après la clôture.
        / The file is a real PDF named after the closure.
        """
        contenu_en_octets, nom_du_fichier, type_du_contenu = generer_pdf_cloture(
            self.cloture
        )

        assert contenu_en_octets[:5] == b"%PDF-"
        assert type_du_contenu == "application/pdf"
        assert nom_du_fichier.startswith(f"cloture-{self.cloture.numero_sequentiel}-")

    # ------------------------------------------------------------------
    # Email de clôture
    # / Closure email
    # ------------------------------------------------------------------

    def test_email_de_cloture_montre_le_chiffre_d_affaires_ttc(self):
        """
        L'email de clôture montre le chiffre d'affaires TTC (17,50 €) sous son vrai
        libellé, et le nombre de ventes.
        / The closure email shows the revenue incl. tax under its true label.
        """
        configuration = Configuration.get_solo()
        configuration.rapport_emails = "tresorerie@example.com"
        configuration.rapport_periodicite = ClotureCaisse.NIVEAU_JOURNALIER
        configuration.save()

        with patch("comptabilite.tasks.CeleryMailerClass") as faux_envoi_de_mail:
            faux_envoi_de_mail.return_value = MagicMock()
            comptabilite.tasks.envoyer_email_cloture(
                self.tenant.schema_name, str(self.cloture.uuid)
            )
        arguments_du_mail = faux_envoi_de_mail.call_args.kwargs

        with translation.override("fr"):
            corps_du_mail = render_to_string(
                arguments_du_mail["template"], arguments_du_mail["context"]
            )

        assert "Chiffre d'affaires TTC" in corps_du_mail
        assert euros("17,50") in corps_du_mail
        assert "Total TTC" not in corps_du_mail

    # ------------------------------------------------------------------
    # FEC : la forme du fichier (contenu : fiche F §4)
    # / FEC: file shape (content: sheet F §4)
    # ------------------------------------------------------------------

    def test_export_fec_a_la_forme_d_un_fec(self):
        """
        Le FEC est un fichier texte UTF-8 sans BOM, tabulé, à 18 colonnes, et ses
        écritures sont équilibrées (Σ débits = Σ crédits) sur le scénario mêlé. Le jus
        est d'abord rangé dans une catégorie de caisse reliée au 707000 : sans compte,
        le FEC serait refusé (fiche E §3.3). Le FEC lit le plan du moment de l'export.
        / The FEC is a UTF-8 (no BOM) tab-separated 18-column file, balanced on the
        mixed scenario (the juice first gets a category with an account).
        """
        from comptabilite.fec import generer_fec_cloture

        categorie_des_boissons = CategorieProduct.objects.create(
            name=f"Boissons exports {identifiant_unique()}",
            compte_comptable=CompteComptable.objects.get(numero_de_compte="707000"),
        )
        produit_du_jus = self.tarif_du_jus.productsold.product
        produit_du_jus.categorie_pos = categorie_des_boissons
        produit_du_jus.save()

        contenu_en_octets, nom_du_fichier, type_du_contenu = generer_fec_cloture(
            self.cloture
        )

        assert type_du_contenu == "text/plain; charset=utf-8"
        assert nom_du_fichier.startswith("FEC-")
        assert not contenu_en_octets.startswith(b"\xef\xbb\xbf")
        lignes_du_fichier = contenu_en_octets.decode("utf-8", errors="strict").split(
            "\r\n"
        )
        premiere_ligne = lignes_du_fichier[0]
        assert premiere_ligne.count("\t") == 17
        assert "JournalCode" in premiere_ligne
        assert "CompteNum" in premiere_ligne

        # Colonnes 12 et 13 : débit et crédit (« 12,34 »).
        # / Columns 12 and 13: debit and credit.
        total_des_debits_en_centimes = 0
        total_des_credits_en_centimes = 0
        nombre_de_lignes_d_ecriture = 0
        for ligne_du_fichier in lignes_du_fichier[1:]:
            if ligne_du_fichier == "":
                continue
            valeurs = ligne_du_fichier.split("\t")
            assert len(valeurs) == 18
            # Un montant du FEC : des chiffres, une VIRGULE, deux chiffres.
            # / A FEC amount: digits, a COMMA, two digits.
            assert re.fullmatch(r"\d+,\d{2}", valeurs[11]), valeurs[11]
            assert re.fullmatch(r"\d+,\d{2}", valeurs[12]), valeurs[12]
            total_des_debits_en_centimes += int(
                Decimal(valeurs[11].replace(",", ".")) * 100
            )
            total_des_credits_en_centimes += int(
                Decimal(valeurs[12].replace(",", ".")) * 100
            )
            nombre_de_lignes_d_ecriture += 1
        assert nombre_de_lignes_d_ecriture > 0
        assert total_des_debits_en_centimes == total_des_credits_en_centimes

    # ------------------------------------------------------------------
    # Mentions légales, chaîne des clôtures
    # / Legal mentions, chain of closures
    # ------------------------------------------------------------------

    def test_pdf_imprime_les_mentions_legales_du_lieu(self):
        """
        Le PDF imprime les mentions légales figées dans le Z : adresse, code postal,
        ville, SIREN, numéro de TVA.
        / The PDF prints the legal mentions stored in the Z.
        """
        html = self._html_du_pdf()

        assert "7 place du Marché" in html
        assert "69001" in html
        assert "Lyon" in html
        assert "552100554" in html
        assert "FR40552100554" in html

    def test_tableur_et_csv_portent_les_mentions_legales(self):
        """
        Les mentions légales sont aussi dans le tableur et le CSV (SIREN).
        / Legal mentions are also in the spreadsheet and the CSV.
        """
        ligne_du_siren_dans_le_csv = self._ligne_qui_commence_par(
            self._lignes_du_csv(), "SIREN / SIRET"
        )
        assert ligne_du_siren_dans_le_csv[1] == "552100554"
        ligne_du_siren_dans_le_tableur = self._valeurs_de_la_ligne_du_tableur(
            self._feuille_du_tableur(), "SIREN / SIRET"
        )
        assert ligne_du_siren_dans_le_tableur[1] == "552100554"

    def test_chaine_des_clotures_saine_apres_le_scenario_mele(self):
        """
        Contrôle de bout en bout : le rapport riche du scénario, stocké en `jsonb` puis
        relu, garde son empreinte (l'aller-retour ne change pas le contenu signé).
        / End-to-end check: the rich report survives the jsonb round trip.
        """
        cle_du_lieu = LaboutikConfiguration.get_solo().get_or_create_hmac_key()

        assert verifier_chaine_clotures(cle_du_lieu) == []

    # ------------------------------------------------------------------
    # Ordre d'affichage stable (le jsonb ne garde pas l'ordre des clés)
    # / Stable display order (jsonb does not keep key order)
    # ------------------------------------------------------------------

    def test_ordre_des_produits_et_des_evenements_stable_apres_la_base(self):
        """
        Le rapport est stocké en `jsonb`, qui range les clés à sa façon (pour des uuid :
        dans l'ordre des uuid). L'affichage trie donc lui-même : produits par nom,
        événements par date. Deux produits dont l'ordre des uuid est l'inverse de
        l'ordre des noms (« Abricot » a le plus grand uuid), deux événements dont
        l'ordre des uuid est l'inverse de l'ordre des dates (« Zénith », le plus tôt,
        a le plus grand uuid). La J est relue depuis la base. Fiche, PDF, tableur et
        CSV : Abricot avant Banane, Zénith avant Atelier.
        / Products sorted by name, events by date, in the four readers, after a
        database round trip.
        """
        moment_des_ventes = timezone.now() + timedelta(minutes=5)
        tarif_vendu_de_l_abricot = self._tarif_vendu_d_un_produit(
            "Abricot ordre",
            uuid.UUID("ffffffff-ffff-4fff-bfff-fffffffffff1"),
            Product.NONE,
        )
        tarif_vendu_de_la_banane = self._tarif_vendu_d_un_produit(
            "Banane ordre",
            uuid.UUID("00000000-0000-4000-8000-000000000001"),
            Product.NONE,
        )
        self._vendre_en_especes_a(moment_des_ventes, tarif_vendu_de_l_abricot)
        self._vendre_en_especes_a(moment_des_ventes, tarif_vendu_de_la_banane)
        self._billet_d_un_evenement_vendu_a(
            moment_des_ventes,
            "Zénith ordre",
            timezone.now() + timedelta(days=1),
            uuid.UUID("ffffffff-ffff-4fff-bfff-fffffffffff2"),
        )
        self._billet_d_un_evenement_vendu_a(
            moment_des_ventes,
            "Atelier ordre",
            timezone.now() + timedelta(days=5),
            uuid.UUID("00000000-0000-4000-8000-000000000002"),
        )
        cloture = self._cloturer_a(moment_des_ventes + timedelta(minutes=5))

        textes_html = {
            "fiche": self._page_de_la_fiche(cloture),
            "PDF": self._html_du_pdf(cloture),
        }
        for nom_du_lecteur, html in textes_html.items():
            assert html.index("Abricot ordre") < html.index("Banane ordre"), (
                nom_du_lecteur
            )
            assert html.index("Zénith ordre") < html.index("Atelier ordre"), (
                nom_du_lecteur
            )

        lignes_du_csv = self._lignes_du_csv(cloture)
        lignes_du_tableur = list(
            self._feuille_du_tableur(cloture).iter_rows(values_only=True)
        )
        lignes_par_lecteur = {"CSV": lignes_du_csv, "tableur": lignes_du_tableur}
        for nom_du_lecteur, lignes in lignes_par_lecteur.items():
            position_de_l_abricot = self._position_de_la_ligne_qui_contient(
                lignes, "Abricot ordre"
            )
            position_de_la_banane = self._position_de_la_ligne_qui_contient(
                lignes, "Banane ordre"
            )
            position_du_zenith = self._position_de_la_ligne_qui_contient(
                lignes, "Zénith ordre"
            )
            position_de_l_atelier = self._position_de_la_ligne_qui_contient(
                lignes, "Atelier ordre"
            )
            assert position_de_l_abricot < position_de_la_banane, nom_du_lecteur
            assert position_du_zenith < position_de_l_atelier, nom_du_lecteur

    # ------------------------------------------------------------------
    # Formules dans les exports
    # / Formulas in the exports
    # ------------------------------------------------------------------

    def _cloture_avec_un_produit_nomme_comme_une_formule(self):
        """
        Une J qui contient la vente d'un produit nommé « =1+1 ».
        / A J with the sale of a product named "=1+1".
        """
        moment_des_ventes = timezone.now() + timedelta(minutes=5)
        tarif_vendu_de_la_formule = self._tarif_vendu_d_un_produit(
            "=1+1", uuid.uuid4(), Product.NONE
        )
        self._vendre_en_especes_a(moment_des_ventes, tarif_vendu_de_la_formule)
        return self._cloturer_a(moment_des_ventes + timedelta(minutes=5))

    def test_csv_un_nom_qui_commence_par_egal_est_precede_d_une_apostrophe(self):
        """
        Dans le CSV, un texte qui commence par « = », « + », « - » ou « @ » est
        précédé d'une apostrophe (l'usage habituel) : un tableur qui ouvre le CSV ne le
        calcule pas. Un produit nommé « =1+1 ».
        / In the CSV, a formula-like text gets a leading apostrophe.
        """
        cloture = self._cloture_avec_un_produit_nomme_comme_une_formule()

        lignes_du_csv = self._lignes_du_csv(cloture)

        self._position_de_la_ligne_qui_contient(lignes_du_csv, "'=1+1")
        for ligne in lignes_du_csv:
            assert "=1+1" not in ligne

    def test_tableur_un_nom_qui_commence_par_egal_est_un_texte_jamais_une_formule(
        self,
    ):
        """
        Dans le tableur, le nom « =1+1 » reste tel quel (aucune apostrophe visible),
        dans une case de type texte qui porte le marqueur natif d'Excel
        (`quotePrefix`) : il n'est jamais calculé. Aucune case du tableur n'est une
        formule.
        / In the spreadsheet, "=1+1" stays as is, in a text cell with Excel's native
        quote prefix; no cell is a formula.
        """
        cloture = self._cloture_avec_un_produit_nomme_comme_une_formule()

        feuille = self._feuille_du_tableur(cloture)

        cases_du_nom = []
        for ligne_de_cases in feuille.iter_rows():
            for case in ligne_de_cases:
                assert case.data_type != "f", case.coordinate
                assert case.value != "'=1+1", case.coordinate
                if case.value == "=1+1":
                    cases_du_nom.append(case)
        assert len(cases_du_nom) >= 1
        for case_du_nom in cases_du_nom:
            assert case_du_nom.data_type == "s"
            assert case_du_nom.quotePrefix is True

    # ------------------------------------------------------------------
    # Langue et heure du lieu
    # / Language and venue time
    # ------------------------------------------------------------------

    def test_pdf_porte_la_langue_active(self):
        """
        `<html lang>` du PDF est la langue active (lecteurs d'écran, césure).
        / The PDF's <html lang> is the active language.
        """
        with translation.override("en"):
            html_en_anglais = html_du_pdf_de_la_cloture(self.cloture)

        assert '<html lang="en">' in html_en_anglais
        assert '<html lang="fr">' in self._html_du_pdf()

    def test_email_de_cloture_en_heure_du_lieu(self):
        """
        Les heures de la période dans l'email sont en heure du lieu (Martinique,
        UTC−4), pas en heure du serveur.
        / The email's period hours are in the venue's time zone.
        """
        fuseau_de_la_martinique = ZoneInfo("America/Martinique")
        configuration = Configuration.get_solo()
        configuration.fuseau_horaire = "America/Martinique"
        configuration.rapport_emails = "tresorerie@example.com"
        configuration.rapport_periodicite = ClotureCaisse.NIVEAU_JOURNALIER
        configuration.save()

        with patch("comptabilite.tasks.CeleryMailerClass") as faux_envoi_de_mail:
            faux_envoi_de_mail.return_value = MagicMock()
            comptabilite.tasks.envoyer_email_cloture(
                self.tenant.schema_name, str(self.cloture.uuid)
            )
        arguments_du_mail = faux_envoi_de_mail.call_args.kwargs
        with translation.override("fr"):
            corps_du_mail = render_to_string(
                arguments_du_mail["template"], arguments_du_mail["context"]
            )

        debut_en_martinique = self.cloture.datetime_debut.astimezone(
            fuseau_de_la_martinique
        ).strftime("%d/%m/%Y %H:%M")
        fin_en_martinique = self.cloture.datetime_fin.astimezone(
            fuseau_de_la_martinique
        ).strftime("%d/%m/%Y %H:%M")
        assert debut_en_martinique in corps_du_mail
        assert fin_en_martinique in corps_du_mail

    def test_nom_des_fichiers_date_en_heure_du_lieu(self):
        """
        Le nom des fichiers exportés est daté du jour de fin de la clôture EN HEURE DU
        LIEU : une J close à 2 h 40 UTC est, en Martinique (UTC−4), la veille à 22 h 40.
        / Export file names are dated in the venue's local time.
        """
        configuration = Configuration.get_solo()
        configuration.fuseau_horaire = "America/Martinique"
        configuration.save()
        apres_demain = (timezone.now() + timedelta(days=2)).date()
        moment_de_la_vente = datetime.combine(
            apres_demain, time(2, 30), tzinfo=ZoneInfo("UTC")
        )
        tarif_vendu = self._tarif_vendu_d_un_produit(
            f"Jus de nuit {identifiant_unique()}", uuid.uuid4(), Product.NONE
        )
        self._vendre_en_especes_a(moment_de_la_vente, tarif_vendu)
        cloture = self._cloturer_a(moment_de_la_vente + timedelta(minutes=10))

        veille_en_martinique = apres_demain - timedelta(days=1)
        date_attendue = veille_en_martinique.strftime("%Y%m%d")
        nom_attendu = f"cloture-{cloture.numero_sequentiel}-{date_attendue}"
        _contenu, nom_du_csv, _type = generer_csv_cloture(cloture)
        _contenu, nom_du_tableur, _type = generer_excel_cloture(cloture)
        _contenu, nom_du_pdf, _type = generer_pdf_cloture(cloture)

        assert nom_du_csv == f"{nom_attendu}.csv"
        assert nom_du_tableur == f"{nom_attendu}.xlsx"
        assert nom_du_pdf == f"{nom_attendu}.pdf"
